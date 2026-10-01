"""Product catalog: required fields, Excel template, upload preview and import."""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.engine import Engine
from sqlalchemy.orm import Session

from app.db.models import Product
from app.db.session import session_scope
from app.i18n import ar, get_language, set_language
from app.services import auth, products
from app.services.seed import seed_defaults

PASSWORD = "StrongPass123"


@pytest.fixture
def session(memory_engine: Engine) -> Iterator[Session]:
    with session_scope(memory_engine) as scoped:
        seed_defaults(scoped)
        yield scoped


@pytest.fixture
def owner_id(session: Session) -> int:
    owner = auth.create_first_owner(
        session, username="owner", password=PASSWORD, password_confirmation=PASSWORD
    )
    return owner.id


@pytest.fixture
def seller_id(session: Session, owner_id: int) -> int:
    seller = auth.create_seller(
        session, owner_id, username="seller", password=PASSWORD, password_confirmation=PASSWORD
    )
    return seller.id


def _sheet(path: Path, rows: list[list[object]], *, title: str = "Sheet") -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = title
    for row in rows:
        worksheet.append(row)
    workbook.save(path)
    return path


def _by_name(session: Session) -> dict[str, Product]:
    return {product.name: product for product in session.scalars(select(Product))}


# ----------------------------------------------------------- required fields
def test_product_needs_only_name_and_price(session: Session, owner_id: int) -> None:
    product = products.create_product(session, owner_id, name="  Redmi 13  ", cash_price=32_000)

    assert (product.name, product.cash_price, product.wholesale_price) == ("Redmi 13", 32_000, 0)


@pytest.mark.parametrize("name", ["", "   "])
def test_product_name_is_required(session: Session, owner_id: int, name: str) -> None:
    with pytest.raises(ValueError):
        products.create_product(session, owner_id, name=name, cash_price=1_000)


@pytest.mark.parametrize("price", [0, -5])
def test_product_price_is_required(session: Session, owner_id: int, price: int) -> None:
    with pytest.raises(ValueError):
        products.create_product(session, owner_id, name="Redmi", cash_price=price)


def test_update_keeps_wholesale_when_not_given(session: Session, owner_id: int) -> None:
    product = products.create_product(session, owner_id, name="A15", cash_price=30_000, wholesale_price=25_000)

    products.update_product(session, owner_id, product.id, name="A15", cash_price=31_000)

    assert (product.cash_price, product.wholesale_price) == (31_000, 25_000)


def test_duplicate_name_has_its_own_error(session: Session, owner_id: int) -> None:
    products.create_product(session, owner_id, name="A15", cash_price=30_000)
    with pytest.raises(products.ProductNameTakenError):
        products.create_product(session, owner_id, name="A15", cash_price=31_000)
    session.rollback()  # the app's session_scope does this on any error


# ------------------------------------------------------------------ template
def test_template_marks_name_and_price_required_and_reads_back(tmp_path: Path, session: Session) -> None:
    path = products.write_product_template(tmp_path / "template.xlsx")

    workbook = load_workbook(path)
    sheet = workbook.worksheets[0]
    headers = [cell.value for cell in sheet[1]]
    assert headers[0] == f"{ar.PROD_TPL_NAME} *"
    assert headers[1] == f"{ar.PROD_TPL_PRICE} *"
    assert ar.PROD_TPL_OPTIONAL in headers[2] and "*" not in headers[2]
    assert sheet.sheet_view.rightToLeft is True
    assert len(workbook.worksheets) == 2  # products + instructions

    sheet.append(["iPhone 13", 110_000, 90_000])
    sheet.append(["Redmi Note 13", 32_000, None])
    workbook.save(path)
    preview = products.preview_product_workbook(session, path)
    assert [(row.name, row.cash_price, row.wholesale_price) for row in preview.rows] == [
        ("iPhone 13", 110_000, 90_000),
        ("Redmi Note 13", 32_000, None),
    ]
    assert preview.rows[0].row_number == 2


@pytest.mark.parametrize("language", ["en", "fr"])
def test_template_follows_the_interface_language(tmp_path: Path, session: Session, language: str) -> None:
    previous = get_language()
    set_language(language)
    try:
        path = products.write_product_template(tmp_path / "template.xlsx")
        expected_header = f"{ar.PROD_TPL_NAME} *"
    finally:
        set_language(previous)
    workbook = load_workbook(path)
    sheet = workbook.worksheets[0]
    assert sheet.sheet_view.rightToLeft is False
    assert sheet["A1"].value == expected_header
    sheet.append(["Galaxy A05", 21_000])
    workbook.save(path)

    # Read back while the interface is in Arabic: headers of every language are recognised.
    assert products.preview_product_workbook(session, path).rows[0].cash_price == 21_000


# ------------------------------------------------------------------- preview
def test_preview_flags_missing_and_invalid_values(tmp_path: Path, session: Session) -> None:
    path = _sheet(tmp_path / "p.xlsx", [
        ["اسم المنتج", "السعر", "سعر الجملة"],
        ["Oppo A18", "٢٥٬٠٠٠ دج", "20,000 DA"],
        [None, 10_000, None],
        ["No price", None, 5_000],
        ["Zero", 0, None],
        ["Text price", "غالي", None],
        ["Bad wholesale", 9_000, "?"],
        [None, None, None],
        ["oppo  a18", 26_000, None],
    ])

    preview = products.preview_product_workbook(session, path)
    rows = {row.row_number: row for row in preview.rows}

    assert rows[2].valid and (rows[2].cash_price, rows[2].wholesale_price) == (25_000, 20_000)
    assert rows[3].errors == (products.ERROR_NAME_MISSING,)
    assert rows[4].errors == (products.ERROR_PRICE_MISSING,)
    assert rows[5].errors == (products.ERROR_PRICE_INVALID,)
    assert rows[6].errors == (products.ERROR_PRICE_INVALID,)
    assert rows[7].errors == (products.ERROR_WHOLESALE_INVALID,)
    assert 8 not in rows  # blank rows are ignored
    assert rows[9].errors == (products.ERROR_DUPLICATE,)
    assert preview.count(products.ACTION_CREATE) == 1
    assert preview.count(products.ACTION_INVALID) == 6


def test_preview_finds_headers_below_a_title_and_in_other_languages(tmp_path: Path, session: Session) -> None:
    path = _sheet(tmp_path / "p.xlsx", [
        ["Liste des prix"],
        [],
        ["Prix de vente", "Désignation", "Prix de gros"],
        [15_000.0, "Nokia 105", 11_000],
    ])

    preview = products.preview_product_workbook(session, path)

    row = preview.rows[0]
    assert (row.row_number, row.name, row.cash_price, row.wholesale_price) == (4, "Nokia 105", 15_000, 11_000)


def test_preview_warns_without_wholesale_or_below_cost(tmp_path: Path, session: Session) -> None:
    path = _sheet(tmp_path / "p.xlsx", [
        ["Product", "Price", "Wholesale"],
        ["Cable", 800, None],
        ["Loss", 1_000, 1_500],
    ])

    rows = products.preview_product_workbook(session, path).rows

    assert rows[0].valid and rows[0].warnings == (products.WARNING_NO_WHOLESALE,)
    assert rows[1].valid and rows[1].warnings == (products.WARNING_BELOW_WHOLESALE,)


def test_preview_matches_existing_products(tmp_path: Path, session: Session, owner_id: int) -> None:
    kept = products.create_product(session, owner_id, name="iPhone 13", cash_price=110_000, wholesale_price=90_000)
    archived = products.create_product(session, owner_id, name="Old phone", cash_price=5_000)
    products.set_product_active(session, owner_id, archived.id, False)
    path = _sheet(tmp_path / "p.xlsx", [
        ["Name", "Price"],
        ["IPHONE 13", 110_000],
        ["Old Phone", 4_000],
        ["Galaxy", 50_000],
    ])

    rows = products.preview_product_workbook(session, path).rows

    assert [(row.action, row.existing_id) for row in rows] == [
        (products.ACTION_UNCHANGED, kept.id),
        (products.ACTION_RESTORE, archived.id),
        (products.ACTION_CREATE, None),
    ]


@pytest.mark.parametrize("rows, code", [
    ([["Customer", "Phone"], ["Ali", "0550"]], "no_header"),
    ([["Name", "Price"]], "empty"),
])
def test_preview_rejects_unusable_sheets(tmp_path: Path, session: Session, rows, code: str) -> None:
    path = _sheet(tmp_path / "p.xlsx", rows)
    with pytest.raises(products.ProductWorkbookError) as error:
        products.preview_product_workbook(session, path)
    assert error.value.code == code


def test_preview_rejects_a_file_that_is_not_excel(tmp_path: Path, session: Session) -> None:
    path = tmp_path / "notes.xlsx"
    path.write_text("not a workbook", encoding="utf-8")
    with pytest.raises(products.ProductWorkbookError) as error:
        products.preview_product_workbook(session, path)
    assert error.value.code == "unreadable"


# -------------------------------------------------------------------- import
def test_import_creates_updates_and_restores(tmp_path: Path, session: Session, owner_id: int) -> None:
    products.create_product(session, owner_id, name="iPhone 13", cash_price=110_000, wholesale_price=90_000)
    archived = products.create_product(session, owner_id, name="Old phone", cash_price=5_000, wholesale_price=3_000)
    products.set_product_active(session, owner_id, archived.id, False)
    path = _sheet(tmp_path / "p.xlsx", [
        ["اسم المنتج *", "السعر *", "سعر الجملة (اختياري)"],
        ["iphone 13", 115_000, None],
        ["Old phone", 4_500, None],
        ["Redmi 13", 32_000, None],
        ["Pixel", 70_000, 60_000],
        ["Broken", None, None],
    ])
    preview = products.preview_product_workbook(session, path)

    result = products.import_products(session, owner_id, preview.rows)

    assert (result.created, result.updated, result.restored, result.unchanged, result.skipped) == (2, 1, 1, 0, 1)
    catalog = _by_name(session)
    assert (catalog["iPhone 13"].cash_price, catalog["iPhone 13"].wholesale_price) == (115_000, 90_000)
    assert catalog["Old phone"].active and catalog["Old phone"].cash_price == 4_500
    assert catalog["Redmi 13"].wholesale_price == 0
    assert catalog["Pixel"].wholesale_price == 60_000
    assert "Broken" not in catalog

    again = products.import_products(session, owner_id, products.preview_product_workbook(session, path).rows)
    assert (again.created, again.updated, again.restored, again.unchanged) == (0, 0, 0, 4)


def test_only_the_owner_can_import(tmp_path: Path, session: Session, seller_id: int) -> None:
    path = _sheet(tmp_path / "p.xlsx", [["Name", "Price"], ["Redmi", 30_000]])
    preview = products.preview_product_workbook(session, path)
    with pytest.raises(auth.AuthorizationError):
        products.import_products(session, seller_id, preview.rows)
    assert _by_name(session) == {}
