"""Products and client-type pages, new client + purchase, history, settings, 2FA sign-in."""

from __future__ import annotations

from datetime import date, datetime, timezone
from pathlib import Path

import pytest
from PySide6.QtCore import QSettings
from PySide6.QtWidgets import QApplication, QMessageBox
from sqlalchemy import select

from app.db.models import Customer, Sale, StockItem, User
from app.db.session import dispose_database_engines, session_scope
from app.i18n import ar
from app.services import auth, backup, categories, customers, payments, products, sales, two_factor


@pytest.fixture(scope="module")
def qapp() -> QApplication:
    return QApplication.instance() or QApplication([])


@pytest.fixture
def shop(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, qapp) -> dict:
    """Migrated file database with an owner, a seller and a small catalog."""
    monkeypatch.setenv("AKREMMOBILE_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.delenv("AKREMMOBILE_DATABASE_URL", raising=False)
    QSettings.setPath(QSettings.Format.NativeFormat, QSettings.Scope.UserScope, str(tmp_path / "settings"))
    for name in ("information", "warning", "critical"):
        monkeypatch.setattr(QMessageBox, name, lambda *args, **kwargs: QMessageBox.StandardButton.Ok)
    from app.db.migrate import upgrade_database
    from app.services.seed import seed_defaults

    upgrade_database()
    seed_defaults()
    with session_scope() as session:
        owner = auth.create_first_owner(
            session, username="akrem", password="Secret123", password_confirmation="Secret123"
        )
        seller = auth.create_seller(
            session, owner.id, username="vendeur", password="Secret123", password_confirmation="Secret123"
        )
        products.create_product(session, owner.id, name="iPhone 13", wholesale_price=90000, cash_price=110000)
        products.create_product(session, owner.id, name="Redmi 13", wholesale_price=25000, cash_price=32000)
        ids = {"owner": owner.id, "seller": seller.id, "tmp": str(tmp_path)}
    yield ids
    dispose_database_engines()


def _plain(text: str) -> str:
    """Cell text without the left-to-right isolate marks used for Latin names in RTL tables."""
    return text.replace("\u2066", "").replace("\u2069", "")


def _user(user_id: int) -> User:
    with session_scope() as session:
        return session.get(User, user_id)


# ------------------------------------------------------------- sidebar pages
def test_products_and_client_types_are_sidebar_pages(shop) -> None:
    from app.ui.main_window import PAGE_CLIENT_TYPES, PAGE_PRODUCTS, MainWindow

    window = MainWindow(current_user=_user(shop["owner"]))
    window._select_page(PAGE_PRODUCTS)
    assert window.pages.currentWidget() is window.products_page
    assert window.topbar.page_title_label.text() == ar.NAV_PRODUCTS
    window._select_page(PAGE_CLIENT_TYPES)
    assert window.pages.currentWidget() is window.client_types_page

    teachers = next(t for t in window.client_types_page.panel._types if t.name == "أساتذة")
    window.client_types_page.view_customers_requested.emit(teachers.id)
    assert window.pages.currentWidget() is window.customers_page
    assert window.customers_page.type_filter.selected_type_id() == teachers.id
    window._quitting = True
    window.close()

    seller_window = MainWindow(current_user=_user(shop["seller"]))
    seller_window._select_page(PAGE_CLIENT_TYPES)
    assert seller_window.pages.currentIndex() != PAGE_CLIENT_TYPES
    seller_window._select_page(PAGE_PRODUCTS)
    assert seller_window.pages.currentIndex() == PAGE_PRODUCTS
    seller_window._quitting = True
    seller_window.close()


def test_products_page_lists_prices_quote_and_hides_wholesale_from_sellers(shop) -> None:
    from app.ui.pages.products_page import ProductDialog, ProductsPage

    page = ProductsPage(_user(shop["owner"]))
    headers = [page.table.horizontalHeaderItem(c).text() for c in range(page.table.columnCount())]
    assert ar.SET_PRODUCT_WHOLESALE in headers
    names = [_plain(page.table.item(r, 0).text()) for r in range(page.table.rowCount())]
    assert names == ["iPhone 13", "Redmi 13"]
    quote_column = page._columns.index("quote")
    # 110,000 + 35% = 148,500 over 6 months
    assert page.table.item(0, quote_column).text().startswith("24,750")

    dialog = ProductDialog(title="x", parent=page)
    dialog.name.setText("Galaxy A15")
    dialog.wholesale_price.setValue(20000)
    dialog.cash_price.setValue(26000)
    assert page.save_dialog(dialog)
    assert page.table.rowCount() == 3

    duplicate = ProductDialog(title="x", parent=page)
    duplicate.name.setText("Galaxy A15")
    duplicate.cash_price.setValue(27000)
    assert not page.save_dialog(duplicate)
    assert duplicate.error.text.text() == ar.PROD_DUPLICATE_NAME

    seller_page = ProductsPage(_user(shop["seller"]))
    seller_headers = [seller_page.table.horizontalHeaderItem(c).text()
                      for c in range(seller_page.table.columnCount())]
    assert ar.SET_PRODUCT_WHOLESALE not in seller_headers
    for button in (seller_page.add_button, seller_page.template_button, seller_page.upload_button):
        assert button.isHidden()


def test_product_dialog_requires_only_name_and_price(shop) -> None:
    from app.ui.pages.products_page import ProductDialog, ProductsPage

    page = ProductsPage(_user(shop["owner"]))
    dialog = ProductDialog(title="x", parent=page)
    assert dialog.wholesale_price.specialValueText() == ar.PROD_OPTIONAL

    assert not page.save_dialog(dialog)
    assert dialog.error.text.text() == ar.PROD_NAME_REQUIRED
    dialog.name.setText("Nokia 105")
    assert not page.save_dialog(dialog)
    assert dialog.error.text.text() == ar.PROD_PRICE_REQUIRED

    dialog.cash_price.setValue(4500)  # wholesale left empty
    assert page.save_dialog(dialog)
    row = [_plain(page.table.item(r, 0).text()) for r in range(page.table.rowCount())].index("Nokia 105")
    assert page.table.item(row, page._columns.index("wholesale")).text() == "—"
    assert page.table.item(row, page._columns.index("margin")).text() == "—"
    with session_scope() as session:
        saved = {p.name: p for p in products.list_products(session)}["Nokia 105"]
        assert (saved.cash_price, saved.wholesale_price) == (4500, 0)


def test_products_page_downloads_template_and_uploads_products(shop, monkeypatch) -> None:
    from openpyxl import load_workbook
    from PySide6.QtWidgets import QFileDialog

    from app.ui.pages import products_page
    from app.ui.pages.products_page import ProductImportDialog, ProductsPage

    page = ProductsPage(_user(shop["owner"]))
    assert page.template_button.isVisibleTo(page) and page.upload_button.isVisibleTo(page)

    target = Path(shop["tmp"]) / "out" / "products"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(target), ""))
    page.template_button.click()
    template = target.with_suffix(".xlsx")
    assert template.exists()

    # The template is the shop's stock sheet: product, wholesale, cash price, battery, colour...
    workbook = load_workbook(template)
    sheet = workbook.worksheets[0]
    sheet.append(["Redmi 13", None, 33000, 0.9, "Blue"])   # one more phone of an existing product
    sheet.append(["Galaxy A05", None, "21,000 دج", "*"])    # new product, new phone, no wholesale
    sheet.append(["Bad row", 1000, None])                  # rejected: no price
    workbook.save(template)

    shown: list[ProductImportDialog] = []

    def fake_exec(dialog):
        shown.append(dialog)
        assert dialog.import_button.isEnabled()
        dialog.import_button.click()
        return 1

    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *a, **k: (str(template), ""))
    monkeypatch.setattr(ProductImportDialog, "exec", fake_exec)
    page.upload_button.click()

    dialog = shown[0]
    assert dialog.preview.is_stock
    results = [dialog.cell_text(r, "result") for r in range(dialog.table.rowCount())]
    assert results == [ar.PROD_ACTION_EXISTING, ar.PROD_ACTION_CREATE, ar.PROD_ACTION_INVALID]
    assert [dialog.cell_text(r, "stock") for r in range(2)] == [ar.PROD_UNIT_ADD.format(count=1)] * 2
    assert dialog.cell_text(0, "battery") == "90%" and dialog.cell_text(1, "battery") == ar.STOCK_BATTERY_NEW
    assert ar.PROD_ERR_PRICE_MISSING in dialog.cell_text(2, "notes")
    assert dialog.result() == 1
    from app.services import stock

    with session_scope() as session:
        catalog = {p.name: p for p in products.list_products(session)}
        # Stock rows never change catalog prices; each phone keeps its own price.
        assert (catalog["Redmi 13"].cash_price, catalog["Redmi 13"].wholesale_price) == (32000, 25000)
        assert (catalog["Galaxy A05"].cash_price, catalog["Galaxy A05"].wholesale_price) == (21000, 0)
        assert "Bad row" not in catalog
        units = {unit.product.name: unit for unit in stock.list_units(session)}
        assert (units["Redmi 13"].cash_price, units["Redmi 13"].battery_health, units["Redmi 13"].color) == (33000, 90, "Blue")
        assert units["Galaxy A05"].is_new
    assert page.stock_panel.table.rowCount() == 2
    names = [_plain(page.table.item(r, 0).text()) for r in range(page.table.rowCount())]
    assert "Galaxy A05" in names

    # A sheet without product columns is refused with a clear message.
    other = Path(shop["tmp"]) / "customers.xlsx"
    from openpyxl import Workbook
    book = Workbook()
    book.active.append(["الاسم واللقب", "الهاتف"])
    book.active.append(["علي", "0550"])
    book.save(other)
    warnings: list[str] = []
    monkeypatch.setattr(products_page.QMessageBox, "warning", lambda _p, _t, text: warnings.append(text))
    assert page.preview_upload(other) is None
    assert warnings == [ar.PROD_UPLOAD_NO_HEADER]


# ----------------------------------------------------- new client + purchase
def test_new_client_form_fills_the_catalog_price_and_saves_the_purchase(shop) -> None:
    from app.ui.pages.customers_page import CustomersPage

    page = CustomersPage(_user(shop["owner"]))
    section = page._purchase_section()
    assert section.isChecked()
    section.product.setCurrentIndex(section.product.findData("iPhone 13"))
    assert section.cash_price.value() == 110000
    section.sale_type.setCurrentIndex(section.sale_type.findData("installment"))
    section.months.setValue(6)
    assert section.rate.value() == 35  # preset for 6 months
    section.down_payment.setValue(28500)
    result = section.calculation()
    assert result is not None and result.total == 148500 and result.monthly_list[0] == 20000
    assert "20,000" in section.summary.text()

    fallback_id = page._types[-1].id
    values = {"full_name": "زبون جديد", "phone": "0550000001", "category_id": fallback_id}
    customer_id = page.save_new_customer(values, section.sale_values())
    assert customer_id is not None
    with session_scope() as session:
        sale = session.scalar(select(Sale).where(Sale.customer_id == customer_id))
        assert (sale.product, sale.sale_type, sale.total, sale.months) == ("iPhone 13", "installment", 148500, 6)

    cash = page._purchase_section()
    cash.product.setCurrentIndex(cash.product.findData("Redmi 13"))
    cash.sale_type.setCurrentIndex(cash.sale_type.findData("cash"))
    cash_id = page.save_new_customer(
        {"full_name": "زبون كاش", "phone": "", "category_id": fallback_id}, cash.sale_values()
    )
    with session_scope() as session:
        assert session.scalar(select(Sale.sale_type).where(Sale.customer_id == cash_id)) == "cash"

    unchecked = page._purchase_section()
    unchecked.setChecked(False)
    assert unchecked.sale_values() is None


# ------------------------------------------------------------------- history
def test_clicking_a_client_shows_their_full_history(shop) -> None:
    from app.ui.dialogs.customer_details_dialog import CustomerDetailsDialog

    with session_scope() as session:
        fallback = categories.get_fallback_type(session)
        customer, sale = sales.create_customer_with_sale(
            session, shop["owner"],
            customer={"full_name": "زبون قديم", "category_id": fallback.id},
            sale={"product": "iPhone 13", "sale_type": "installment", "wholesale_price": 90000,
                  "cash_price": 110000, "rate": 35, "down_payment": 28500, "months": 6,
                  "purchase_date": date(2026, 1, 10)},
        )
        sales.create_sale(session, customer_id=customer.id, product="Redmi 13", sale_type="cash",
                          wholesale_price=25000, cash_price=32000, purchase_date=date(2026, 3, 2))
        first = min(sale.installments, key=lambda row: row.installment_index)
        payments.record_installment_payment(session, first.id, 20000, date(2026, 2, 3), "CCP")
        customer_id = customer.id

    dialog = CustomerDetailsDialog(customer_id, _user(shop["owner"]))
    assert dialog.tabs.count() == 2
    summary = dialog.history_summary
    assert summary.values["count"].text() == "2"
    assert summary.values["bought"].text() == f"{148500 + 32000:,} دج"
    assert summary.values["paid"].text() == f"{28500 + 20000 + 32000:,} دج"
    assert summary.profile.text() == ar.HIST_PROFILE_FACILITIES

    panel = dialog.history_panel
    assert panel.purchases_table.rowCount() == 2
    assert _plain(panel.purchases_table.item(0, 1).text()) == "Redmi 13"  # newest first
    assert panel.purchases_table.item(0, 2).text() == ar.SALE_CASH
    assert panel.payments_table.rowCount() == 1
    assert panel.payments_table.item(0, 2).text() == "CCP"

    panel.purchase_activated.emit(dialog.sale_selector.itemData(1))
    assert dialog.tabs.currentIndex() == 0


# ------------------------------------------------------------------ settings
def test_settings_can_move_the_backup_folder(shop) -> None:
    from app.ui.pages.settings_page import SettingsPage

    page = SettingsPage(_user(shop["owner"]))
    target = Path(shop["tmp"]) / "usb" / "akrem-backups"
    assert page.set_backup_folder(str(target))
    assert page.backup_folder_label.text() == str(target.resolve())
    created = backup.create_backup()
    assert created.parent == target.resolve()
    assert page.set_backup_folder(None)
    assert backup.backup_directory() == backup.default_backup_directory()

    a_file = Path(shop["tmp"]) / "not-a-folder"
    a_file.write_text("x")
    assert not page.set_backup_folder(str(a_file))


def test_owner_can_add_owner_and_seller_accounts(shop) -> None:
    from app.ui.pages.settings_page import SettingsPage, _UserCredentialsDialog

    page = SettingsPage(_user(shop["owner"]))
    roles = [page.seller_table.item(r, 1).text() for r in range(page.seller_table.rowCount())]
    assert ar.ROLE_OWNER in roles and ar.ROLE_SELLER in roles

    original = _UserCredentialsDialog.exec

    def fake_exec(dialog):
        dialog.username.setText("partner")
        dialog.password.setText("Partner123")
        dialog.confirmation.setText("Partner123")
        dialog.role.setCurrentIndex(dialog.role.findData("owner"))
        return 1

    _UserCredentialsDialog.exec = fake_exec
    try:
        page._create_seller()
    finally:
        _UserCredentialsDialog.exec = original
    with session_scope() as session:
        partner = session.scalar(select(User).where(User.username == "partner"))
        assert partner.role == "owner"
    assert auth.authenticate(
        session_scope().__enter__(), username="partner", password="Partner123"
    ).status == "success"


def test_password_change_from_settings(shop) -> None:
    from app.ui.pages.settings_page import SettingsPage

    page = SettingsPage(_user(shop["seller"]))
    page.current_password.setText("Secret123")
    page.new_password.setText("Changed456")
    page.new_password_confirmation.setText("Changed456")
    page._change_own_password()
    with session_scope() as session:
        assert auth.authenticate(session, username="vendeur", password="Changed456").status == "success"


# ------------------------------------------------- Google Authenticator 2FA
def test_two_factor_setup_shows_a_qr_code_and_protects_sign_in(shop, monkeypatch) -> None:
    from app.ui.dialogs.auth_dialogs import LoginDialog
    from app.ui.dialogs.two_factor_dialog import TwoFactorSetupDialog
    from app.ui.pages.settings_page import SettingsPage

    owner = _user(shop["owner"])
    page = SettingsPage(owner)
    assert page.tfa_status.text() == ar.TFA_STATUS_OFF

    setup = TwoFactorSetupDialog(owner, page)
    assert setup.qr.pixmap() is not None and not setup.qr.pixmap().isNull()
    assert setup.key_label.text().replace(" ", "") == setup.secret
    setup.code.setText("000000")
    if two_factor.totp_code(setup.secret, datetime.now(timezone.utc)) != "000000":
        assert not setup.confirm()
        assert setup.error.text.text() == ar.TFA_INVALID_CODE
    setup.code.setText(two_factor.totp_code(setup.secret, datetime.now(timezone.utc)))
    assert setup.confirm()
    page._refresh_two_factor()
    assert page.tfa_status.text() == ar.TFA_STATUS_ON

    login = LoginDialog()
    login.username.setText("akrem")
    login.password.setText("Secret123")
    login._login()
    assert login.user is None and login.code_page.isVisibleTo(login)
    login.code.setText("999999" if setup.secret else "")
    login._verify_code()
    assert login.user is None
    # A fresh code for the next 30-second window (the current one was used to enable).
    future = datetime.fromtimestamp(datetime.now(timezone.utc).timestamp() + 30, timezone.utc)
    login.code.setText(two_factor.totp_code(setup.secret, future))
    login._verify_code()
    assert login.user is not None and login.user.username == "akrem"

    page._reset_user_two_factor  # owner reset is available
    with session_scope() as session:
        auth.reset_two_factor(session, shop["owner"], shop["owner"])
    page._refresh_two_factor()
    assert page.tfa_status.text() == ar.TFA_STATUS_OFF


# ------------------------------------------------------- flexible payment plans
def test_product_dialog_saves_a_default_plan_shown_on_the_products_page(shop) -> None:
    from app.ui.pages.products_page import ProductDialog, ProductsPage

    page = ProductsPage(_user(shop["owner"]))
    dialog = ProductDialog(title="x", parent=page)
    assert (dialog.plan_months.value(), dialog.plan_interval.value(), dialog.plan_rate.value()) == (6, 1, -1)
    assert dialog.plan_interval.text() == ar.PLAN_EVERY_1
    dialog.name.setText("Pixel 8")
    dialog.cash_price.setValue(100000)
    dialog.plan_months.setValue(10)
    dialog.plan_interval.setValue(2)
    assert page.save_dialog(dialog)

    with session_scope() as session:
        saved = {p.name: p for p in products.list_products(session)}["Pixel 8"]
        assert products.plan_of(saved) == products.ProductPlan(10, 2, None)
    row = [_plain(page.table.item(r, 0).text()) for r in range(page.table.rowCount())].index("Pixel 8")
    plan_cell = page.table.item(row, page._columns.index("plan")).text()
    assert "10" in plan_cell and ar.PLAN_EVERY_2 in plan_cell and "40%" in plan_cell
    # 100,000 + 40% = 140,000 in 5 payments every 2 months
    quote = page.table.item(row, page._columns.index("quote")).text()
    assert quote.startswith("28,000") and ar.PLAN_EVERY_2 in quote and quote.endswith("× 5")

    # Shrinking the plan below the interval caps the interval.
    edit = ProductDialog(title="x", parent=page, product=saved)
    edit.plan_months.setValue(1)
    assert edit.plan().interval == 1


def test_new_client_purchase_starts_from_the_product_plan_and_can_differ(shop) -> None:
    from app.ui.pages.customers_page import CustomersPage

    with session_scope() as session:
        product = {p.name: p for p in products.list_products(session)}["Redmi 13"]
        products.update_product(
            session, shop["owner"], product.id, name="Redmi 13", cash_price=32000,
            plan=products.ProductPlan(months=4, interval=2, rate=20),
        )
    page = CustomersPage(_user(shop["owner"]))
    section = page._purchase_section()
    section.product.setCurrentIndex(section.product.findData("Redmi 13"))
    section.sale_type.setCurrentIndex(section.sale_type.findData("installment"))
    assert (section.months.value(), section.interval.value(), section.rate.value()) == (4, 2, 20)
    assert ar.PLAN_EVERY_2 in section.plan_hint.text()
    assert "19,200" in section.summary.text()  # 38,400 in 2 payments

    # This client pays monthly over 5 months instead (rate follows the preset).
    section.months.setValue(5)
    section.interval.setValue(1)
    assert section.rate.value() == 35
    values = section.sale_values()
    assert (values["months"], values["payment_interval"]) == (5, 1)

    section.months.setValue(6)
    section.interval.setValue(3)
    fallback_id = page._types[-1].id
    customer_id = page.save_new_customer(
        {"full_name": "زبون كل 3 أشهر", "phone": "0550000009", "category_id": fallback_id},
        section.sale_values(),
    )
    with session_scope() as session:
        sale = session.scalar(select(Sale).where(Sale.customer_id == customer_id))
        assert (sale.months, sale.payment_interval, len(sale.installments)) == (6, 3, 2)


def test_new_sale_page_applies_the_product_plan_and_saves_the_interval(shop, monkeypatch) -> None:
    from app.ui.pages.new_sale_page import NewSalePage

    with session_scope() as session:
        product = {p.name: p for p in products.list_products(session)}["iPhone 13"]
        products.update_product(
            session, shop["owner"], product.id, name="iPhone 13", cash_price=110000,
            plan=products.ProductPlan(months=12, interval=6),
        )
        customer = customers.create_customer(
            session, full_name="زبون الصفحة", phone="0550000123", category_id=categories.get_fallback_type(session).id,
        )
        customer_id = customer.id
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: None)
    page = NewSalePage(current_user=_user(shop["owner"]))
    page.sale_type.setCurrentIndex(page.sale_type.findData("installment"))
    page.product.setCurrentText("iPhone 13")
    assert (page.months.value(), page.interval.value(), page.rate.value()) == (12, 6, 45)
    assert page.monthly_caption.text() == ar.INSTALLMENT_AMOUNT
    assert page.schedule_table.rowCount() == 2  # 2 payments, every 6 months

    page.interval.setValue(1)  # this client pays monthly instead
    assert page.schedule_table.rowCount() == 12
    assert page.monthly_caption.text() == ar.MONTHLY_AMOUNT
    page.interval.setValue(4)
    page._customer_id = customer_id
    page._save_sale()
    with session_scope() as session:
        sale = session.scalar(select(Sale).where(Sale.customer_id == customer_id))
        assert (sale.months, sale.payment_interval, len(sale.installments)) == (12, 4, 3)


# ------------------------------------------------------------ stock (phones)
def test_stock_tab_adds_edits_filters_and_removes_phones(shop, monkeypatch) -> None:
    from app.services import stock
    from app.ui.pages.products_page import ProductsPage

    page = ProductsPage(_user(shop["owner"]))
    panel = page.stock_panel
    assert page.tabs.currentWidget() is panel and panel.empty_card.isVisibleTo(panel)

    dialog = panel.new_dialog()
    dialog.product.setCurrentText("iphone 13")  # existing name, other case
    assert dialog.product_state.text() == ar.SALE_PRODUCT_EXISTS
    assert dialog.cash_price.value() == 110000  # catalog price filled in
    dialog.battery.setValue(91)
    dialog.color.setCurrentText("Blue")
    dialog.imei.setText("35678901234567")
    dialog.quantity.setValue(1)
    assert panel.save_dialog(dialog)

    dialog = panel.new_dialog()
    dialog.product.setCurrentText("Galaxy S24")  # not in the catalog yet
    assert dialog.product_state.text() == ar.SALE_PRODUCT_NEW
    assert not panel.save_dialog(dialog)
    assert dialog.error.text.text() == ar.PROD_PRICE_REQUIRED
    dialog.cash_price.setValue(150000)
    dialog.is_new.setChecked(True)
    dialog.quantity.setValue(2)
    assert panel.save_dialog(dialog)

    panel.refresh()
    assert panel.table.rowCount() == 3
    assert "3" in page.count_label.text()
    column = panel._columns.index
    assert _plain(panel.table.item(0, column("product")).text()) == "iPhone 13"
    assert panel.table.item(0, column("battery")).text() == "91%"
    assert panel.table.item(1, column("battery")).text() == ar.STOCK_BATTERY_NEW

    panel.condition_filter.setCurrentIndex(panel.condition_filter.findData(stock.CONDITION_NEW))
    assert panel.table.rowCount() == 2
    panel.condition_filter.setCurrentIndex(0)
    panel.color_filter.setCurrentIndex(panel.color_filter.findData("Blue"))
    assert panel.table.rowCount() == 1
    panel.color_filter.setCurrentIndex(0)
    panel.search.setText("4567")
    assert panel.table.rowCount() == 1
    panel.search.clear()

    with session_scope() as session:
        assert products.find_by_name(session, "galaxy s24").cash_price == 150000
        unit = stock.list_units(session)[0]
    edit = panel.new_dialog(unit)
    assert edit.imei.text() == "35678901234567" and not edit._form.isRowVisible(edit.quantity)
    edit.note.setText("small scratch")
    assert panel.save_dialog(edit, unit.id)

    panel.table.setCurrentCell(2, 0)
    assert panel.delete_button.isEnabled()
    assert panel.delete_unit(panel.selected_unit().id)
    assert panel.table.rowCount() == 2

    exported = page.export_stock(Path(shop["tmp"]) / "stock-out")
    from openpyxl import load_workbook

    assert load_workbook(exported).worksheets[0].max_row == 3

    seller_page = ProductsPage(_user(shop["seller"]))
    assert seller_page.stock_panel.add_button.isHidden()
    assert "wholesale" not in seller_page.stock_panel._columns


def test_new_sale_page_sells_a_phone_from_stock_with_its_details(shop, monkeypatch) -> None:
    from app.services import stock
    from app.ui.pages.new_sale_page import NewSalePage

    with session_scope() as session:
        product = products.find_by_name(session, "Redmi 13")
        (unit,) = stock.add_units(
            session, shop["owner"], product_id=product.id, cash_price=33000, wholesale_price=26000,
            details=stock.UnitDetails(color="Black", battery_health=92),
        )
        unit_id, reference = unit.id, unit.reference
        customer = customers.create_customer(
            session, full_name="مشتري المخزون", phone="0550000444",
            category_id=categories.get_fallback_type(session).id,
        )
        customer_id = customer.id
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: None)
    page = NewSalePage(current_user=_user(shop["owner"]))
    page.sale_type.setCurrentIndex(page.sale_type.findData("installment"))
    page.product.setCurrentText("Redmi 13")
    details = page.sale_details
    assert details.product_state.text() == ar.SALE_PRODUCT_EXISTS
    assert details.unit.count() == 2  # "none" + the phone in stock
    details.unit.setCurrentIndex(details.unit.findData(unit_id))
    assert page.cash_price.value() == 33000 and page.wholesale_price.value() == 26000
    assert details.reference.text() == reference and details.reference.isReadOnly()
    page._customer_id = customer_id
    page._save_sale()
    with session_scope() as session:
        sale = session.scalar(select(Sale).where(Sale.customer_id == customer_id))
        assert (sale.cash_price, sale.color, sale.battery_health, sale.reference) == (33000, "Black", 92, reference)
        sold = session.get(StockItem, unit_id)
        assert sold.status == stock.STATUS_SOLD and sold.sale_id == sale.id

    # The cash module records the same details and links a brand-new product name.
    page.cash_product.setCurrentText("Tablet Modio")
    assert page.cash_details.product_state.text() == ar.SALE_PRODUCT_NEW
    page.cash_total.setValue(21000)
    page.cash_details.color.setCurrentText("Grey")
    page.cash_details.is_new.setChecked(True)
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.StandardButton.No)
    page._save_cash_sale()
    with session_scope() as session:
        cash_sale = session.scalar(select(Sale).where(Sale.product == "Tablet Modio"))
        assert (cash_sale.color, cash_sale.is_new) == ("Grey", True)
        assert cash_sale.product_id == products.find_by_name(session, "tablet modio").id


def test_new_client_form_links_a_typed_product_and_records_phone_details(shop) -> None:
    from app.ui.pages.customers_page import CustomersPage

    page = CustomersPage(_user(shop["owner"]))
    section = page._purchase_section()
    assert section.product.isEditable()
    section.product.setCurrentText("Honor X9D 12/256")
    assert section.details.product_state.text() == ar.SALE_PRODUCT_NEW
    section.sale_type.setCurrentIndex(section.sale_type.findData("cash"))
    section.cash_price.setValue(82800)
    section.details.battery.setValue(97)
    section.details.color.setCurrentText("Black")
    values = section.sale_values()
    assert (values["product"], values["battery_health"], values["color"]) == ("Honor X9D 12/256", 97, "Black")
    customer_id = page.save_new_customer(
        {"full_name": "زبون هونر", "phone": "0550000555", "category_id": page._types[-1].id}, values
    )
    with session_scope() as session:
        sale = session.scalar(select(Sale).where(Sale.customer_id == customer_id))
        assert sale.product_id == products.find_by_name(session, "honor x9d 12/256").id
        assert (sale.battery_health, sale.color) == (97, "Black")


def test_clients_page_has_purchase_and_end_of_deduction_filters(shop) -> None:
    from app.ui.pages.customers_page import CustomersPage

    with session_scope() as session:
        fallback = categories.get_fallback_type(session).id
        recent = customers.create_customer(session, full_name="اشترى حديثاً", category_id=fallback)
        sales.create_sale(session, customer_id=recent.id, product="X", sale_type="installment",
                          wholesale_price=0, cash_price=6000, months=6, purchase_date=date.today())
        old = customers.create_customer(session, full_name="قديم", category_id=fallback)
        sales.create_sale(session, customer_id=old.id, product="X", sale_type="installment",
                          wholesale_price=0, cash_price=6000, months=3, purchase_date=date(2024, 1, 10))
    page = CustomersPage(_user(shop["owner"]))
    page.purchased_filter.setCurrentIndex(page.purchased_filter.findData("this_month"))
    assert "اشترى حديثاً" in _visible_names(page) and "قديم" not in _visible_names(page)
    page.purchased_filter.setCurrentIndex(0)
    page.deduction_filter.setCurrentIndex(page.deduction_filter.findData("finished"))
    assert _visible_names(page) == {"قديم"}
    page._reset_filters()
    assert page.deduction_filter.currentIndex() == 0


def _visible_names(page) -> set[str]:
    """Customer names currently listed on the clients page."""
    model = page.table.model()
    return {model.index(row, 1).data() for row in range(model.rowCount())}


# ------------------------------------------------------------- debt book pages
def test_debt_pages_add_pay_by_facility_and_show_the_schedule(shop, monkeypatch) -> None:
    from app.services import debts
    from app.ui.main_window import PAGE_DEBTS_IN, PAGE_DEBTS_OUT, MainWindow
    from app.ui.pages.debts_page import DebtDetailsDialog, PaymentDialog, ReceivablesPage

    window = MainWindow(current_user=_user(shop["owner"]))
    window._select_page(PAGE_DEBTS_IN)
    page = window.pages.currentWidget()
    assert isinstance(page, ReceivablesPage) and page.empty_card.isVisibleTo(page)

    dialog = page.new_dialog()
    assert not page.save_dialog(dialog)
    assert dialog.error.text.text() == ar.DEBT_ERR_NAME
    dialog.first_name.setText("علي")
    dialog.last_name.setText("بن صالح")
    dialog.email.setText("not-an-email")
    dialog.amount.setValue(60000)
    assert not page.save_dialog(dialog)
    assert dialog.error.text.text() == ar.DEBT_ERR_EMAIL
    dialog.email.setText("ali@example.com")
    dialog.national_id.setText("109876543210987654")
    dialog.plan.setCurrentIndex(dialog.plan.findData(debts.PLAN_INSTALLMENTS))
    dialog.months.setValue(6)
    dialog.interval.setValue(2)
    assert ar.PLAN_EVERY_2 in dialog.preview.text() and "20,000" in dialog.preview.text()
    assert page.save_dialog(dialog)

    assert page.table.rowCount() == 1
    assert page.table.item(0, 0).text() == "علي بن صالح"
    assert page.cards["remaining"].text() == "60,000 دج"
    page.table.setCurrentCell(0, 0)
    summary = page.selected()
    assert summary.plan == debts.PLAN_INSTALLMENTS and page.pay_button.isEnabled()

    def pay(dialog: PaymentDialog) -> int:
        dialog.amount.setValue(20000)
        dialog.save_button.click()
        return 1

    monkeypatch.setattr(PaymentDialog, "exec", pay)
    assert page.record_payment_for(summary)
    assert page.cards["paid"].text() == "20,000 دج"

    details = page.open_details()
    assert isinstance(details, DebtDetailsDialog)
    assert details.schedule_table.rowCount() == 3 and details.payments_table.rowCount() == 1
    assert details.schedule_table.item(0, 4).text() == ar.DEBT_STATUS_PAID
    details.close()

    window._select_page(PAGE_DEBTS_OUT)
    payables = window.pages.currentWidget()
    assert payables.direction == debts.PAYABLE and payables.table.rowCount() == 0
    seller = MainWindow(current_user=_user(shop["seller"]))
    seller._select_page(PAGE_DEBTS_IN)
    assert seller.pages.currentIndex() != PAGE_DEBTS_IN  # owner only


def test_dashboard_shows_the_net_financial_position(shop) -> None:
    from app.services import debts
    from app.services.debts import DebtTerms, PersonDetails
    from app.ui.pages.dashboard_page import DashboardPage

    with session_scope() as session:
        debts.create_debt(session, shop["owner"], direction=debts.RECEIVABLE,
                          person=PersonDetails("سمير", "ب"), terms=DebtTerms(30000, date.today()))
        debts.create_debt(session, shop["owner"], direction=debts.PAYABLE,
                          person=PersonDetails("مورد", "الهواتف"), terms=DebtTerms(50000, date.today()))
    page = DashboardPage(_user(shop["owner"]))
    assert page.position_box.isVisibleTo(page)
    assert page.debtors_owe_card.value_label.text() == "30,000 دج"
    assert page.shop_owes_card.value_label.text() == "50,000 دج"
    assert page.net_value.text() == f"{30000 - 50000:,} دج"
    opened: list[str] = []
    page.debts_requested.connect(opened.append)
    page.shop_owes_card.clicked.emit()
    assert opened == ["payable"]
    assert not DashboardPage(_user(shop["seller"])).position_box.isVisibleTo(page)
