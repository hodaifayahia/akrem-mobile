"""Stock units, the products & stock sheet, sales linked to products/stock, new filters.

Workbooks are synthetic but laid out like the shop's "Gros et Détail" stock
sheet and its installment sheet (repeated header rows, ``*`` for new phones,
Excel percentages, notes typed in odd columns).
"""

from __future__ import annotations

from collections.abc import Iterator
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.engine import Engine
from sqlalchemy.orm import Session

from app.db.models import Category, Customer, Product, Sale, StockItem
from app.db.session import session_scope
from app.services import auth, customers, importer, product_sheet, products, sales, stock
from app.services.seed import seed_defaults

PASSWORD = "StrongPass123"
TODAY = date(2026, 10, 15)


@pytest.fixture
def session(memory_engine: Engine) -> Iterator[Session]:
    with session_scope(memory_engine) as scoped:
        seed_defaults(scoped)
        yield scoped


@pytest.fixture
def owner_id(session: Session) -> int:
    return auth.create_first_owner(
        session, username="owner", password=PASSWORD, password_confirmation=PASSWORD
    ).id


@pytest.fixture
def seller_id(session: Session, owner_id: int) -> int:
    return auth.create_seller(
        session, owner_id, username="seller", password=PASSWORD, password_confirmation=PASSWORD
    ).id


@pytest.fixture
def customer_id(session: Session) -> int:
    category = session.scalar(select(Category).limit(1))
    return customers.create_customer(session, full_name="زبون تجريبي", category_id=category.id).id


def _workbook(path: Path, rows: list[list[object]]) -> Path:
    book = Workbook()
    for row in rows:
        book.active.append(row)
    book.save(path)
    return path


SHOP_HEADER = ["المنتج", "سعر الجملة", "سعر الكاش", "البطارية", "اللون", "SATUTS", "REF"]


def _shop_sheet(path: Path) -> Path:
    """A stock sheet in the shop's layout: batches separated by repeated headers."""
    return _workbook(path, [
        SHOP_HEADER,
        ["Tab Lenovo 3/64", 19000, 22000, "*", "Black", None, None],
        ["Phone A 4/64", 22500, 26000, "*", "Black", "+++++++++++", "REF-00001"],
        ["Phone A 4/64", 22500, 26000, "*", "Black", "+++++++++++", "REF-00001"],
        ["Phone B", 26000, 30500, 0.92, "Red", "+++++++++++", "REF-00002"],
        ["Phone B", 29000, 33500, 1, "Black", None, "REF-00003"],
        [None, None, None, None, None, None, None],
        ["المنتج", "سعر الجملة", "سعر الكاش", "البطارية", "اللون", "IMEI", "REF"],
        ["Phone C 6/128", 31000, 34000, None, None, None, None],
        ["Phone C 6/128", 31000, 34000, None, None, None, None],
        ["Phone D 4/128", 30000, 30000, "كريم ب.", None, None, None],
        ["Phone E", 20000, 22500, None, None, "Sold to Sami at cost", None],
        ["Phone F", None, None, "*", "Orange", None, None],
        ["Phone G 8/256", 41000, 45000, "95%", "Blue", "356789012345678", None],
    ])


# ------------------------------------------------------------------ parsing
@pytest.mark.parametrize("value, expected", [
    ("*", (True, None, None)),
    ("جديد", (True, None, None)),
    (0.92, (False, 92, None)),
    (1, (False, 100, None)),
    (93, (False, 93, None)),
    ("88%", (False, 88, None)),
    ("٨٥٪", (False, 85, None)),
    (None, (False, None, None)),
    ("خطأ", (False, None, "خطأ")),
])
def test_battery_cells(value, expected) -> None:
    assert stock.parse_battery(value) == expected


@pytest.mark.parametrize("value, expected", [
    ("356789012345678", "356789012345678"),
    (356789012345678.0, "356789012345678"),
    ("35-678901-234567-8", "356789012345678"),
    ("+++++++++++", None),
    ("1234", None),
])
def test_imei_cells(value, expected) -> None:
    assert stock.normalize_imei(value) == expected


# ------------------------------------------------------------------- service
def test_units_get_automatic_references_and_unique_imeis(session: Session, owner_id: int) -> None:
    product = products.create_product(session, owner_id, name="Phone A", cash_price=26000, wholesale_price=22500)
    first, second = stock.add_units(
        session, owner_id, product_id=product.id, quantity=2,
        details=stock.UnitDetails(color="Black", is_new=True),
    )
    assert (first.reference, second.reference) == ("REF-00001", "REF-00002")
    assert (first.cash_price, first.wholesale_price, first.is_new) == (26000, 22500, True)
    (with_imei,) = stock.add_units(
        session, owner_id, product_id=product.id, cash_price=25000,
        details=stock.UnitDetails(imei="356789012345678", battery_health=92),
    )
    assert with_imei.reference == "REF-00003" and with_imei.cash_price == 25000
    with pytest.raises(stock.StockError) as taken:
        stock.add_units(session, owner_id, product_id=product.id, details=stock.UnitDetails(imei="356789012345678"))
    assert taken.value.code == "imei_taken"
    with pytest.raises(ValueError):
        stock.add_units(session, owner_id, product_id=product.id, quantity=2,
                        details=stock.UnitDetails(reference="REF-9"))
    assert stock.stock_counts(session) == {product.id: (3, 0)}


def test_only_the_owner_changes_stock(session: Session, owner_id: int, seller_id: int) -> None:
    product = products.create_product(session, owner_id, name="Phone A", cash_price=26000)
    with pytest.raises(auth.AuthorizationError):
        stock.add_units(session, seller_id, product_id=product.id)


def test_stock_filters(session: Session, owner_id: int) -> None:
    a = products.create_product(session, owner_id, name="Phone A", cash_price=26000)
    b = products.create_product(session, owner_id, name="Phone B", cash_price=30000)
    stock.add_units(session, owner_id, product_id=a.id, details=stock.UnitDetails(color="Black", is_new=True))
    stock.add_units(session, owner_id, product_id=b.id, details=stock.UnitDetails(color="Red", battery_health=80))
    stock.add_units(session, owner_id, product_id=b.id, details=stock.UnitDetails(color="black", battery_health=97))
    stock.add_units(session, owner_id, product_id=b.id, details=stock.UnitDetails(note="screen scratch"))

    def names(**filters) -> list[str]:
        return [f"{unit.product.name}/{unit.color}" for unit in stock.list_units(session, **filters)]

    assert names(condition=stock.CONDITION_NEW) == ["Phone A/Black"]
    assert names(condition=stock.CONDITION_USED) == ["Phone B/Red", "Phone B/black"]
    assert names(condition=stock.CONDITION_LOW_BATTERY) == ["Phone B/Red"]
    assert names(condition=stock.CONDITION_UNKNOWN) == ["Phone B/None"]
    assert names(color="BLACK") == ["Phone A/Black", "Phone B/black"]
    assert names(product_id=a.id) == ["Phone A/Black"]
    assert names(search="scratch") == ["Phone B/None"]
    assert names(search="ref-00002") == ["Phone B/Red"]
    assert stock.colors(session) == ["Black", "black", "Red"] or stock.colors(session) == ["black", "Black", "Red"]


def test_sold_units_cannot_be_deleted(session: Session, owner_id: int, customer_id: int) -> None:
    product = products.create_product(session, owner_id, name="Phone A", cash_price=26000)
    (unit,) = stock.add_units(session, owner_id, product_id=product.id)
    sales.create_sale_for_user(
        session, owner_id, customer_id=customer_id, product="Phone A", sale_type="cash",
        wholesale_price=0, cash_price=26000, stock_item_id=unit.id,
    )
    with pytest.raises(stock.StockError):
        stock.delete_unit(session, owner_id, unit.id)


# ---------------------------------------------------------------- the sheet
def test_shop_stock_sheet_imports_every_phone_once(tmp_path: Path, session: Session, owner_id: int) -> None:
    path = _shop_sheet(tmp_path / "gros.xlsx")
    preview = product_sheet.preview_product_workbook(session, path)

    assert preview.is_stock
    rows = {row.row_number: row for row in preview.rows}
    assert 8 not in rows and 7 not in rows  # blank row and the repeated header
    assert (rows[2].is_new, rows[2].color, rows[2].cash_price, rows[2].wholesale_price) == (True, "Black", 22000, 19000)
    assert (rows[5].battery_health, rows[6].battery_health) == (92, 100)
    assert rows[3].reference == "REF-00001" and rows[3].note == "+++++++++++"
    assert rows[4].reference is None and product_sheet.WARNING_REF_DUPLICATE in rows[4].warnings
    assert rows[11].note == "كريم ب." and product_sheet.WARNING_BATTERY_TEXT in rows[11].warnings
    assert rows[12].note == "Sold to Sami at cost" and product_sheet.WARNING_STATUS_TEXT in rows[12].warnings
    assert rows[13].errors == (product_sheet.ERROR_PRICE_MISSING,)
    assert rows[14].imei == "356789012345678" and rows[14].battery_health == 95
    assert preview.units_to_add == 10
    assert preview.count(product_sheet.ACTION_CREATE) == 7  # Tab, A, B, C, D, E, G

    result = product_sheet.import_products(session, owner_id, preview.rows)
    assert (result.created, result.units_added, result.skipped) == (7, 10, 1)
    references = sorted(unit.reference for unit in stock.list_units(session))
    assert references[:3] == ["REF-00001", "REF-00002", "REF-00003"]
    assert len(set(references)) == 10
    phone_b = products.find_by_name(session, "phone b")
    assert (phone_b.cash_price, phone_b.wholesale_price) == (30500, 26000)  # first row's prices

    again = product_sheet.preview_product_workbook(session, path)
    assert again.units_to_add == 0 and not again.has_changes


def test_reference_match_updates_the_unit_and_sold_units_are_left_alone(
    tmp_path: Path, session: Session, owner_id: int, customer_id: int
) -> None:
    product_sheet.import_products(
        session, owner_id, product_sheet.preview_product_workbook(session, _shop_sheet(tmp_path / "a.xlsx")).rows
    )
    unit = stock.find_available_unit(session, reference="REF-00002")
    sales.create_sale_for_user(
        session, owner_id, customer_id=customer_id, product="Phone B", sale_type="cash",
        wholesale_price=0, cash_price=1, stock_item_id=unit.id,
    )
    path = _workbook(tmp_path / "b.xlsx", [
        SHOP_HEADER,
        ["Phone B", 26000, 31000, 0.9, "Red", None, "REF-00002"],
        ["Phone B", 29000, 32000, 1, "Black", None, "REF-00003"],
    ])
    rows = product_sheet.preview_product_workbook(session, path).rows
    assert [row.unit_action for row in rows] == [product_sheet.UNIT_SOLD, product_sheet.UNIT_UPDATE]
    product_sheet.import_products(session, owner_id, rows)
    assert stock.find_available_unit(session, reference="REF-00003").cash_price == 32000


def test_template_and_export_follow_the_shop_layout(tmp_path: Path, session: Session, owner_id: int) -> None:
    template = product_sheet.write_product_template(tmp_path / "t.xlsx")
    headers = [cell.value for cell in load_workbook(template).worksheets[0][1]]
    assert headers[0].startswith("المنتج") and headers[0].endswith("*")
    assert headers[1].startswith("سعر الجملة") and headers[2].startswith("سعر الكاش")
    assert headers[3:8] == ["البطارية", "اللون", "IMEI", "REF", "ملاحظات"]

    product_sheet.import_products(
        session, owner_id, product_sheet.preview_product_workbook(session, _shop_sheet(tmp_path / "a.xlsx")).rows
    )
    exported = product_sheet.export_stock_workbook(session, tmp_path / "e.xlsx", stock.list_units(session))
    sheet = load_workbook(exported).worksheets[0]
    assert sheet.max_row == 11
    assert sheet["D2"].value == "*" and sheet["G2"].value.startswith("REF-")
    assert not product_sheet.preview_product_workbook(session, exported).has_changes


def test_price_list_without_stock_columns_keeps_catalog_behaviour(tmp_path: Path, session: Session, owner_id: int) -> None:
    path = _workbook(tmp_path / "p.xlsx", [["المنتج", "السعر"], ["Phone A", 26000], ["Phone A", 27000]])
    preview = product_sheet.preview_product_workbook(session, path)
    assert not preview.is_stock
    assert preview.rows[1].errors == (product_sheet.ERROR_DUPLICATE,)
    product_sheet.import_products(session, owner_id, preview.rows)
    assert session.scalar(select(func.count(StockItem.id))) == 0


# ------------------------------------------------------- sales and products
def test_owner_sale_links_or_creates_the_product(session: Session, owner_id: int, customer_id: int) -> None:
    products.create_product(session, owner_id, name="iPhone 13", cash_price=110000)
    linked = sales.create_sale_for_user(
        session, owner_id, customer_id=customer_id, product="IPHONE  13", sale_type="cash",
        wholesale_price=90000, cash_price=110000,
    )
    created = sales.create_sale_for_user(
        session, owner_id, customer_id=customer_id, product="Redmi 15C", sale_type="cash",
        wholesale_price=30000, cash_price=35000, color="Blue", battery_health=88, imei="356789012345678",
    )
    assert linked.product == "iPhone 13" and linked.product_id is not None
    new_product = products.find_by_name(session, "redmi 15c")
    assert created.product_id == new_product.id and new_product.cash_price == 35000
    assert (created.color, created.battery_health, created.imei) == ("Blue", 88, "356789012345678")


def test_selling_a_stock_unit_copies_its_details_and_price(
    session: Session, owner_id: int, seller_id: int, customer_id: int
) -> None:
    product = products.create_product(session, owner_id, name="Phone B", cash_price=30000, wholesale_price=25000)
    (unit,) = stock.add_units(
        session, owner_id, product_id=product.id, cash_price=33500, wholesale_price=29000,
        details=stock.UnitDetails(color="Black", battery_health=100, imei="356789012345678"),
    )
    sale = sales.create_sale_for_user(
        session, seller_id, customer_id=customer_id, product="Phone B", sale_type="installment",
        wholesale_price=0, cash_price=0, rate=35, months=6, stock_item_id=unit.id,
    )
    assert (sale.cash_price, sale.wholesale_price) == (33500, 29000)  # the unit's own prices
    assert (sale.color, sale.battery_health, sale.imei, sale.reference) == ("Black", 100, "356789012345678", unit.reference)
    assert (unit.status, unit.sale_id, unit.sold_at) == (stock.STATUS_SOLD, sale.id, sale.purchase_date)
    with pytest.raises(stock.StockError):
        sales.create_sale_for_user(
            session, owner_id, customer_id=customer_id, product="Phone B", sale_type="cash",
            wholesale_price=0, cash_price=1, stock_item_id=unit.id,
        )


def test_typing_an_imei_assigns_the_matching_stock_unit(session: Session, owner_id: int, customer_id: int) -> None:
    product = products.create_product(session, owner_id, name="Phone B", cash_price=30000)
    (unit,) = stock.add_units(session, owner_id, product_id=product.id, details=stock.UnitDetails(imei="356789012345678"))
    sale = sales.create_sale(
        session, customer_id=customer_id, product="Phone B", sale_type="cash",
        wholesale_price=0, cash_price=30000, imei="35 6789 0123 45678",
    )
    assert unit.status == stock.STATUS_SOLD and unit.sale_id == sale.id


def test_invalid_imei_on_a_sale_is_refused(session: Session, customer_id: int) -> None:
    with pytest.raises(ValueError):
        sales.create_sale(
            session, customer_id=customer_id, product="X", sale_type="cash",
            wholesale_price=0, cash_price=1000, imei="12AB",
        )


# ------------------------------------------------------ installment sheet
def test_sales_sheet_optional_columns(tmp_path: Path, session: Session, owner_id: int) -> None:
    product = products.create_product(session, owner_id, name="Phone B", cash_price=30000)
    stock.add_units(session, owner_id, product_id=product.id, details=stock.UnitDetails(reference="REF-00042"))
    teachers = session.scalar(select(Category).where(Category.name == "أساتذة"))
    headers = list(importer.HEADERS.values())
    row = {
        "full_name": "زبون أ", "product": "phone b", "sale_type": "بالتقسيط", "wholesale_price": 25000,
        "cash_price": 30000, "rate": 35, "down_payment": 0, "months": 6, "purchase_date": date(2026, 7, 20),
        "phone": 550123456, "client_type": "أساتذة", "payment_interval": 2, "color": "Red",
        "battery": 0.9, "reference": "REF-00042",
    }
    second = dict(row, full_name="زبون ب", product="Brand new model", phone="0661 23 45 67", client_type="نوع غير معروف",
                  payment_interval=None, reference=None, imei="356789012345678", battery="*")
    path = _workbook(tmp_path / "s.xlsx", [headers] + [[r.get(field) for field in importer.HEADERS] for r in (row, second)])
    book = load_workbook(path)
    book.active.title = importer.SHEET_NAME
    book.save(path)

    preview = importer.preview_workbook(path)
    first_row, second_row = preview.rows
    assert (first_row.phone, first_row.client_type, first_row.payment_interval) == ("0550123456", "أساتذة", 2)
    assert (first_row.battery_health, second_row.is_new, second_row.imei) == (90, True, "356789012345678")
    assert len(first_row.calculation.monthly_list) == 3

    fallback = customers.get_fallback_category(session) if hasattr(customers, "get_fallback_category") else None
    from app.services import categories

    result = importer.import_preview(
        session, preview, uncategorized_id=(fallback or categories.get_fallback_type(session)).id,
        mark_past_due_paid=False, owner_user_id=owner_id,
    )
    assert result.imported == 2
    buyer = session.scalar(select(Customer).where(Customer.full_name == "زبون أ"))
    assert (buyer.phone, buyer.category_id) == ("0550123456", teachers.id)
    sale_a, = buyer.sales
    assert (sale_a.product_id, sale_a.payment_interval, len(sale_a.installments)) == (product.id, 2, 3)
    assert stock.find_available_unit(session, reference="REF-00042") is None  # sold through the sheet
    new_product = products.find_by_name(session, "brand new model")
    assert new_product is not None
    other = session.scalar(select(Customer).where(Customer.full_name == "زبون ب"))
    assert other.sales[0].product_id == new_product.id and other.category.name == "غير مصنف"


# ------------------------------------------------------------------ filters
def _sale(session: Session, customer: int, *, bought: date, months: int | None = 6, kind: str = "installment") -> Sale:
    return sales.create_sale(
        session, customer_id=customer, product="X", sale_type=kind, wholesale_price=0, cash_price=6000,
        months=months if kind == "installment" else None, purchase_date=bought,
    )


def test_purchase_period_and_end_of_deduction_filters(session: Session) -> None:
    category = session.scalar(select(Category).limit(1))

    def client(name: str) -> int:
        return customers.create_customer(session, full_name=name, category_id=category.id).id

    recent = client("اشترى هذا الشهر")
    _sale(session, recent, bought=date(2026, 10, 3))
    ending = client("ينتهي هذا الشهر")
    _sale(session, ending, bought=date(2026, 4, 2))  # 6 months -> ends 02/10/2026
    next_month = client("ينتهي الشهر القادم")
    _sale(session, next_month, bought=date(2026, 5, 20))
    finished = client("انتهى")
    _sale(session, finished, bought=date(2025, 1, 5))
    cash_only = client("كاش")
    _sale(session, cash_only, bought=date(2026, 9, 9), kind="cash")

    def names(**filters) -> set[str]:
        return {row.full_name for row in customers.list_customers(
            session, year=2026, month=10, today=TODAY, **filters)}

    assert names(purchased="this_month") == {"اشترى هذا الشهر"}
    assert names(purchased="last_month") == {"كاش"}
    assert names(purchased="last_3_months") == {"اشترى هذا الشهر", "كاش"}
    assert names(purchased="this_year") == {"اشترى هذا الشهر", "ينتهي هذا الشهر", "ينتهي الشهر القادم", "كاش"}
    assert names(deduction="ends_this_month") == {"ينتهي هذا الشهر"}
    assert names(deduction="ends_next_month") == {"ينتهي الشهر القادم"}
    assert names(deduction="finished") == {"انتهى", "ينتهي هذا الشهر"}
    assert names(deduction="ongoing") == {"اشترى هذا الشهر", "ينتهي الشهر القادم"}
    with pytest.raises(ValueError):
        names(deduction="soon")


def test_search_finds_clients_by_imei_or_reference(session: Session, customer_id: int) -> None:
    sales.create_sale(
        session, customer_id=customer_id, product="Phone", sale_type="cash", wholesale_price=0,
        cash_price=1000, imei="356789012345678", reference="REF-00777",
    )
    for text in ("45678", "ref-00777"):
        rows = customers.list_customers(session, year=2026, month=10, today=TODAY, search=text)
        assert [row.id for row in rows] == [customer_id]
