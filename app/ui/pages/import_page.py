"""Excel workbook preview, mapping, validation, and import workflow."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any, Mapping

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QApplication,
    QCheckBox,
    QComboBox,
    QFormLayout,
    QFileDialog,
    QGroupBox,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMessageBox,
    QProgressBar,
    QPushButton,
    QScrollArea,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

from app import config
from app.db.session import session_scope
from app.i18n import ar
from app.services import categories, importer
from app.ui.events import events


def _field_labels() -> dict[str, str]:
    """Labels in the active language (evaluated at call time)."""
    return {
        "full_name": ar.IMP_FIELD_FULL_NAME,
        "product": ar.IMP_FIELD_PRODUCT,
        "sale_type": ar.IMP_FIELD_SALE_TYPE,
        "wholesale_price": ar.IMP_FIELD_WHOLESALE,
        "cash_price": ar.IMP_FIELD_CASH_PRICE,
        "rate": ar.IMP_FIELD_RATE,
        "financed": ar.IMP_FIELD_FINANCED,
        "down_payment": ar.IMP_FIELD_DOWN_PAYMENT,
        "months": ar.IMP_FIELD_MONTHS,
        "monthly_amount": ar.IMP_FIELD_MONTHLY,
        "profit": ar.IMP_FIELD_PROFIT,
        "purchase_date": ar.IMP_FIELD_PURCHASE_DATE,
        "end_date": ar.IMP_FIELD_END_DATE,
        "phone": ar.IMP_FIELD_PHONE,
        "client_type": ar.IMP_FIELD_CLIENT_TYPE,
        "payment_interval": ar.IMP_FIELD_INTERVAL,
        "color": ar.PROD_TPL_COLOR,
        "battery": ar.PROD_TPL_BATTERY,
        "imei": ar.PROD_TPL_IMEI,
        "reference": ar.PROD_TPL_REF,
    }


def _preview_columns() -> tuple[str, ...]:
    """Labels in the active language (evaluated at call time)."""
    return (
        ar.IMP_COLUMN_ROW,
        ar.IMP_COLUMN_CUSTOMER,
        ar.IMP_COLUMN_PRODUCT,
        ar.IMP_COLUMN_TYPE,
        ar.IMP_COLUMN_WHOLESALE,
        ar.IMP_COLUMN_CASH,
        ar.IMP_COLUMN_RATE,
        ar.IMP_COLUMN_DOWN,
        ar.IMP_COLUMN_MONTHS,
        ar.IMP_COLUMN_DATE,
        ar.IMP_COLUMN_SHEET_FINANCED,
        ar.IMP_COLUMN_CALC_FINANCED,
        ar.IMP_COLUMN_SHEET_MONTHLY,
        ar.IMP_COLUMN_CALC_MONTHLY,
        ar.IMP_COLUMN_SHEET_PROFIT,
        ar.IMP_COLUMN_CALC_PROFIT,
        ar.IMP_COLUMN_STATUS,
        ar.IMP_COLUMN_SOURCE,
    )


def _sale_type_choices() -> tuple[tuple[str, str], ...]:
    """Labels in the active language (evaluated at call time)."""
    return (
        (ar.IMP_TYPE_CASH, "cash"),
        (ar.IMP_TYPE_INSTALLMENT, "installment"),
        (ar.IMP_TYPE_CREDIT, "credit"),
    )


class ImportPage(QWidget):
    """Guide an owner through validating and importing an Excel workbook."""

    def __init__(self, current_user: Any, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.current_user = current_user
        self._path: Path | None = None
        self._preview: importer.ImportPreview | None = None
        self._mapping_widgets: dict[str, QComboBox] = {}
        self._sale_type_overrides: dict[int, str] = {}
        self._sheet_value_rows: set[int] = set()
        self._uncategorized_id: int | None = None
        self._loading_preview = False

        if getattr(current_user, "role", None) != "owner":
            layout = QVBoxLayout(self)
            layout.setContentsMargins(34, 30, 34, 30)
            layout.addWidget(QLabel(ar.IMP_OWNER_REQUIRED, self))
            return
        self._build_ui()
        self._load_uncategorized_category()

    def _build_ui(self) -> None:
        """Create file controls, mapping form, preview table, and import actions."""
        layout = QVBoxLayout(self)
        layout.setContentsMargins(28, 22, 28, 22)
        layout.setSpacing(14)

        heading_row = QHBoxLayout()
        title = QLabel(ar.IMP_PAGE_TITLE, self)
        title.setObjectName("pageTitle")
        heading_row.addWidget(title)
        heading_row.addStretch(1)
        self.template_button = QPushButton(f"📄  {ar.IMP_TEMPLATE}", self)
        self.template_button.setProperty("variant", "secondary")
        self.template_button.setCursor(Qt.CursorShape.PointingHandCursor)
        self.template_button.clicked.connect(self._create_template)
        heading_row.addWidget(self.template_button)
        layout.addLayout(heading_row)

        file_row = QHBoxLayout()
        self.file_path = QLineEdit(self)
        self.file_path.setReadOnly(True)
        self.file_path.setPlaceholderText(ar.IMP_FILE)
        self.pick_file_button = QPushButton(f"📁  {ar.IMP_PICK_FILE}", self)
        self.pick_file_button.setProperty("variant", "secondary")
        self.pick_file_button.setCursor(Qt.CursorShape.PointingHandCursor)
        self.pick_file_button.clicked.connect(self._pick_file)
        file_row.addWidget(self.file_path, 1)
        file_row.addWidget(self.pick_file_button)
        layout.addLayout(file_row)

        sheet_row = QHBoxLayout()
        self.sheet_label = QLabel(f"📑  {ar.IMP_SHEET}:", self)
        self.sheet_label.setStyleSheet("color: #9DBEFF; font-weight: 600; font-size: 13px;")
        self.sheet_selector = QComboBox(self)
        self.sheet_selector.setMinimumWidth(240)
        self.sheet_selector.currentIndexChanged.connect(self._sheet_changed)
        sheet_row.addWidget(self.sheet_label)
        sheet_row.addWidget(self.sheet_selector)
        sheet_row.addStretch(1)
        layout.addLayout(sheet_row)

        self.mapping_section = QGroupBox(f"⚙️  {ar.IMP_MAPPING_TITLE}", self)
        mapping_layout = QVBoxLayout(self.mapping_section)
        mapping_layout.addWidget(QLabel(ar.IMP_MAPPING_HELP, self.mapping_section))
        self.mapping_form = QFormLayout()
        mapping_layout.addLayout(self.mapping_form)
        self.mapping_button_row = QHBoxLayout()
        self.mapping_button_row.addStretch(1)
        self.mapping_apply_button = QPushButton(f"✓  {ar.IMP_MAP_APPLY}", self.mapping_section)
        self.mapping_apply_button.setProperty("variant", "primary")
        self.mapping_apply_button.setCursor(Qt.CursorShape.PointingHandCursor)
        self.mapping_apply_button.clicked.connect(self._apply_mapping)
        self.mapping_button_row.addWidget(self.mapping_apply_button)
        mapping_layout.addLayout(self.mapping_button_row)
        self.mapping_section.setVisible(False)
        layout.addWidget(self.mapping_section)

        self.mapping_toggle = QPushButton(f"⚙️  {ar.IMP_MAP_COLUMNS}", self)
        self.mapping_toggle.setProperty("variant", "ghost")
        self.mapping_toggle.setCursor(Qt.CursorShape.PointingHandCursor)
        self.mapping_toggle.clicked.connect(self._toggle_mapping)
        self.mapping_toggle.setVisible(False)
        layout.addWidget(self.mapping_toggle, alignment=Qt.AlignmentFlag.AlignRight)

        self.recalculated_note = QLabel(ar.IMP_RECALCULATED_NOTE, self)
        self.recalculated_note.setStyleSheet("color: #8A94A6; font-size: 13px;")
        self.recalculated_note.setWordWrap(True)
        layout.addWidget(self.recalculated_note)

        category_row = QHBoxLayout()
        cat_lbl = QLabel(f"🏷️  {ar.IMP_CATEGORY}:", self)
        cat_lbl.setStyleSheet("color: #9DBEFF; font-weight: 600; font-size: 13px;")
        category_row.addWidget(cat_lbl)
        self.category_selector = QComboBox(self)
        self.category_selector.setMinimumWidth(240)
        self.category_selector.currentIndexChanged.connect(self._refresh_import_availability)
        category_row.addWidget(self.category_selector)
        self.category_missing = QLabel(ar.IMP_CATEGORY_MISSING, self)
        self.category_missing.setStyleSheet("color: #EF4444; font-size: 12px; font-weight: 600;")
        self.category_missing.setWordWrap(True)
        self.category_missing.setVisible(False)
        category_row.addWidget(self.category_missing, 1)
        layout.addLayout(category_row)

        self.mark_past_due = QCheckBox(ar.IMP_MARK_PAST_DUE, self)
        self.mark_past_due.setChecked(True)
        layout.addWidget(self.mark_past_due)

        self.preview_summary = QLabel(ar.IMP_NO_ROWS, self)
        self.preview_summary.setStyleSheet("color: #9DBEFF; font-weight: 600; font-size: 13px;")
        layout.addWidget(self.preview_summary)

        self.preview_table = QTableWidget(0, len(_preview_columns()), self)
        self.preview_table.setHorizontalHeaderLabels(_preview_columns())
        self.preview_table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.preview_table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.preview_table.setSelectionMode(QTableWidget.SelectionMode.SingleSelection)
        self.preview_table.verticalHeader().setVisible(False)
        self.preview_table.horizontalHeader().setStretchLastSection(True)
        self.preview_table.setAlternatingRowColors(True)
        self.preview_table.setShowGrid(False)
        self.preview_table.setMinimumHeight(270)
        scroll = QScrollArea(self)
        scroll.setWidgetResizable(True)
        scroll.setWidget(self.preview_table)
        layout.addWidget(scroll, 1)

        bottom_row = QHBoxLayout()
        self.progress = QProgressBar(self)
        self.progress.setRange(0, 0)
        self.progress.setFormat(ar.IMP_PROGRESS)
        self.progress.setVisible(False)
        bottom_row.addWidget(self.progress, 1)
        self.import_button = QPushButton(f"📥  {ar.IMP_IMPORT}", self)
        self.import_button.setProperty("variant", "primary")
        self.import_button.setCursor(Qt.CursorShape.PointingHandCursor)
        self.import_button.clicked.connect(self._run_import)
        self.import_button.setEnabled(False)
        bottom_row.addWidget(self.import_button)
        layout.addLayout(bottom_row)

    def _pick_file(self) -> None:
        """Choose a workbook and load its available worksheet names."""
        selected, _filter = QFileDialog.getOpenFileName(
            self,
            ar.IMP_PICK_FILE,
            "",
            ar.IMP_FILE_FILTER,
        )
        if not selected:
            return
        self._path = Path(selected)
        self.file_path.setText(str(self._path))
        self._sale_type_overrides.clear()
        self._sheet_value_rows.clear()
        self.sheet_selector.blockSignals(True)
        self.sheet_selector.clear()
        self.sheet_selector.blockSignals(False)
        try:
            sheet_names = _sheet_names(self._path)
        except Exception:
            self._show_error(ar.IMP_FILE_ERROR)
            self._preview = None
            self._render_preview()
            return

        self.sheet_selector.blockSignals(True)
        self.sheet_selector.clear()
        self.sheet_selector.addItems(sheet_names)
        desired_index = self.sheet_selector.findText(importer.SHEET_NAME)
        self.sheet_selector.setCurrentIndex(desired_index if desired_index >= 0 else 0)
        self.sheet_selector.blockSignals(False)
        self.sheet_selector.setEnabled(bool(sheet_names))
        self.sheet_label.setEnabled(bool(sheet_names))
        self._load_preview()

    def _sheet_changed(self, _index: int) -> None:
        """Reload preview when another worksheet is selected."""
        if self._path is not None and self.sheet_selector.currentText():
            self._sale_type_overrides.clear()
            self._sheet_value_rows.clear()
            self._load_preview()

    def _load_preview(
        self,
        *,
        mapping: Mapping[str, int] | None = None,
        preserve_overrides: bool = False,
    ) -> None:
        """Read the selected workbook with the importer service and refresh UI."""
        if self._path is None:
            return
        overrides = dict(self._sale_type_overrides) if preserve_overrides else {}
        try:
            self._preview = importer.preview_workbook(
                self._path,
                sheet_name=self.sheet_selector.currentText() or importer.SHEET_NAME,
                mapping=mapping,
                sale_type_overrides=overrides,
            )
        except Exception:
            self._preview = None
            self.mapping_section.setVisible(False)
            self.mapping_toggle.setVisible(False)
            self._show_error(ar.IMP_FILE_ERROR)
            self._render_preview()
            return

        self._sale_type_overrides = {
            row.source_row: row.sale_type
            for row in self._preview.rows
            if row.sale_type is not None and row.source_row in overrides
        }
        self._sheet_value_rows.intersection_update(
            row.source_row
            for row in self._preview.rows
            if row.ready and _has_discrepancy(row)
        )
        self._populate_mapping_form(self._preview)
        needs_mapping = bool(self._preview.missing_columns)
        self.mapping_section.setVisible(needs_mapping)
        self.mapping_toggle.setVisible(not needs_mapping)
        if needs_mapping:
            self.mapping_section.setTitle(ar.IMP_MAPPING_TITLE)
        self._render_preview()

    def _populate_mapping_form(self, preview: importer.ImportPreview) -> None:
        """Populate one source-column picker for each known import field."""
        while self.mapping_form.rowCount():
            self.mapping_form.removeRow(0)
        self._mapping_widgets.clear()
        for field in importer.HEADERS:
            combo = QComboBox(self.mapping_section)
            combo.addItem(ar.IMP_COLUMN_NONE, None)
            for column, header in enumerate(preview.headers):
                display_header = header or ar.IMP_EMPTY_HEADER
                combo.addItem(
                    ar.IMP_COLUMN_OPTION.format(
                        column=column + 1,
                        header=display_header,
                    ),
                    column,
                )
            existing_column = preview.mapping.get(field)
            if existing_column is not None:
                selected_index = combo.findData(existing_column)
                if selected_index >= 0:
                    combo.setCurrentIndex(selected_index)
            label = _field_labels()[field]
            if field in importer.REQUIRED_HEADERS:
                label += ar.IMP_REQUIRED_MARK
            self.mapping_form.addRow(label, combo)
            self._mapping_widgets[field] = combo

    def _toggle_mapping(self) -> None:
        """Let the user manually remap columns even when headers matched."""
        if self._preview is None:
            return
        self.mapping_section.setVisible(not self.mapping_section.isVisible())

    def _apply_mapping(self) -> None:
        """Validate required field selections and parse rows with the mapping."""
        mapping: dict[str, int] = {}
        for field, combo in self._mapping_widgets.items():
            column = combo.currentData()
            if column is not None:
                mapping[field] = int(column)
        missing = [field for field in importer.REQUIRED_HEADERS if field not in mapping]
        if missing:
            QMessageBox.warning(self, ar.IMP_MAPPING_TITLE, ar.IMP_MAPPING_INCOMPLETE)
            return
        self._load_preview(mapping=mapping, preserve_overrides=True)

    def _render_preview(self) -> None:
        """Render validated rows, diagnostics, and sheet-versus-calculated values."""
        self._loading_preview = True
        self.preview_table.setRowCount(0)
        if self._preview is None:
            self.preview_summary.setText(ar.IMP_FILE_ERROR if self._path else ar.IMP_NO_ROWS)
            self._loading_preview = False
            self._refresh_import_availability()
            return

        preview = self._preview
        self.preview_table.setRowCount(len(preview.rows))
        if not preview.rows:
            self.preview_summary.setText(ar.IMP_NO_ROWS)
        else:
            ready = sum(1 for row in preview.rows if row.ready)
            warning_count = sum(1 for row in preview.rows if row.warnings)
            error_count = sum(1 for row in preview.rows if row.errors)
            self.preview_summary.setText(
                ar.IMP_PREVIEW_SUMMARY.format(
                    sheet=preview.sheet_name,
                    rows=len(preview.rows),
                    ready=ready,
                    warnings=warning_count,
                    errors=error_count,
                )
            )

        for row_index, row in enumerate(preview.rows):
            calculation = row.calculation
            monthly = calculation.monthly_list[0] if calculation and calculation.monthly_list else None
            values = (
                str(row.source_row),
                row.full_name,
                row.product,
                self._sale_type_label(row.sale_type) if row.sale_type else ar.IMP_CHOOSE_TYPE,
                _format_integer(row.wholesale_price),
                _format_integer(row.cash_price),
                _format_integer(row.rate),
                _format_integer(row.down_payment),
                _format_integer(row.months),
                _format_date(row.purchase_date),
                _format_integer(row.sheet_financed),
                _format_integer(calculation.financed if calculation else None),
                _format_integer(row.sheet_monthly),
                _format_integer(monthly),
                _format_integer(row.sheet_profit),
                _format_integer(calculation.profit if calculation else None),
            )
            for column, value in enumerate(values):
                self._set_cell(row_index, column, value)

            status_text, diagnostics = _row_diagnostics(row)
            status_item = self._set_cell(row_index, 16, status_text)
            status_item.setToolTip(diagnostics)
            if row.errors:
                status_item.setForeground(Qt.GlobalColor.red)
            elif row.warnings:
                status_item.setForeground(Qt.GlobalColor.darkYellow)
            else:
                status_item.setForeground(Qt.GlobalColor.darkGreen)

            if row.sale_type is None:
                type_selector = self._make_sale_type_selector(row.source_row)
                self.preview_table.setCellWidget(row_index, 3, type_selector)
            if row.sheet_financed is not None and _has_discrepancy(row):
                sheet_selector = QCheckBox(ar.IMP_USE_SHEET_VALUES, self.preview_table)
                sheet_selector.setChecked(row.source_row in self._sheet_value_rows)
                sheet_selector.setToolTip(ar.IMP_USE_SHEET_TIP)
                sheet_selector.toggled.connect(
                    lambda selected, source_row=row.source_row:
                    self._set_sheet_value_source(source_row, selected)
                )
                self.preview_table.setCellWidget(row_index, 17, sheet_selector)
            else:
                self._set_cell(row_index, 17, ar.IMP_RECALCULATED)

        self.preview_table.resizeColumnsToContents()
        self.preview_table.setColumnWidth(16, 320)
        self.preview_table.setColumnWidth(17, 210)
        self.preview_table.resizeRowsToContents()
        self._loading_preview = False
        self._refresh_import_availability()

    def _make_sale_type_selector(self, source_row: int) -> QComboBox:
        """Create a per-row selector for missing or unrecognized sale types."""
        combo = QComboBox(self.preview_table)
        combo.addItem(ar.IMP_CHOOSE_TYPE, None)
        for label, value in _sale_type_choices():
            combo.addItem(label, value)
        selected_type = self._sale_type_overrides.get(source_row)
        if selected_type is not None:
            index = combo.findData(selected_type)
            if index >= 0:
                combo.setCurrentIndex(index)
        combo.currentIndexChanged.connect(
            lambda _index, row=source_row, selector=combo:
            self._sale_type_changed(row, selector.currentData())
        )
        return combo

    def _sale_type_changed(self, source_row: int, sale_type: str | None) -> None:
        """Rebuild validation state after the user chooses a missing type."""
        if self._loading_preview or self._preview is None:
            return
        if sale_type is None:
            self._sale_type_overrides.pop(source_row, None)
        else:
            self._sale_type_overrides[source_row] = str(sale_type)
        self._load_preview(
            mapping=dict(self._preview.mapping),
            preserve_overrides=True,
        )

    def _set_sheet_value_source(self, source_row: int, use_sheet: bool) -> None:
        """Remember the selected discrepancy source for one Excel row."""
        if self._loading_preview:
            return
        if use_sheet:
            self._sheet_value_rows.add(source_row)
        else:
            self._sheet_value_rows.discard(source_row)

    def _refresh_import_availability(self, *_args: object) -> None:
        """Enable import only when mapping, category, and blank sale types are resolved."""
        if not hasattr(self, "import_button"):
            return
        preview = self._preview
        can_import = (
            preview is not None
            and not preview.missing_columns
            and self._uncategorized_id is not None
            and bool(preview.rows)
            and any(row.ready for row in preview.rows)
            and all(row.sale_type is not None for row in preview.rows)
        )
        self.import_button.setEnabled(can_import)

    def _load_uncategorized_category(self) -> None:
        """Require the existing default uncategorized category before importing."""
        try:
            with session_scope() as session:
                category = categories.get_fallback_type(session)
                if category is not None:
                    self._uncategorized_id = category.id
                    self.category_selector.clear()
                    self.category_selector.addItem(category.name, category.id)
                    self.category_selector.setEnabled(True)
                    self.category_missing.setVisible(False)
                else:
                    self._uncategorized_id = None
                    self.category_selector.clear()
                    self.category_selector.setEnabled(False)
                    self.category_missing.setVisible(True)
        except Exception:
            self._uncategorized_id = None
            self.category_selector.clear()
            self.category_selector.setEnabled(False)
            self.category_missing.setVisible(True)
        self._refresh_import_availability()

    def _run_import(self) -> None:
        """Import validated rows inside one session transaction and save a report."""
        if self._preview is None or self._uncategorized_id is None:
            self._refresh_import_availability()
            return
        if self._preview.missing_columns:
            QMessageBox.warning(self, ar.IMP_PAGE_TITLE, ar.IMP_MAPPING_INCOMPLETE)
            return
        if any(row.sale_type is None for row in self._preview.rows):
            QMessageBox.warning(self, ar.IMP_PAGE_TITLE, ar.IMP_TYPE_REQUIRED)
            return
        answer = QMessageBox.question(
            self,
            ar.IMP_PAGE_TITLE,
            ar.IMP_CONFIRM_IMPORT,
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No,
        )
        if answer != QMessageBox.StandardButton.Yes:
            return

        self.import_button.setEnabled(False)
        self.progress.setVisible(True)
        QApplication.processEvents()
        try:
            with session_scope() as session:
                result = importer.import_preview(
                    session,
                    self._preview,
                    uncategorized_id=self._uncategorized_id,
                    mark_past_due_paid=self.mark_past_due.isChecked(),
                    use_sheet_values=set(self._sheet_value_rows),
                    owner_user_id=int(self.current_user.id),
                )
        except Exception:
            self.progress.setVisible(False)
            self._refresh_import_availability()
            self._show_error(ar.IMP_IMPORT_ERROR)
            return
        self.progress.setVisible(False)

        report_path: Path | None = None
        try:
            report_path = _save_import_report(self._preview, result)
        except Exception:
            report_path = None
        events.data_changed.emit()
        events.notify.emit(
            "success",
            ar.IMP_DONE_TITLE,
            ar.IMP_DONE_TOAST.format(count=result.imported),
            5000,
        )
        summary = ar.IMP_DONE_SUMMARY.format(
            imported=result.imported,
            skipped=result.skipped,
            warnings=result.warnings,
        )
        if report_path is not None:
            summary += "\n" + ar.IMP_REPORT_PATH.format(path=report_path)
        else:
            summary += "\n" + ar.IMP_REPORT_SAVE_FAILED
        QMessageBox.information(self, ar.IMP_DONE_TITLE, summary)
        self._refresh_import_availability()

    def _create_template(self) -> None:
        """Create an Excel template with headers and common-value dropdowns."""
        selected, _filter = QFileDialog.getSaveFileName(
            self,
            ar.IMP_TEMPLATE_TITLE,
            ar.IMP_TEMPLATE_FILENAME,
            ar.IMP_TEMPLATE_FILTER,
        )
        if not selected:
            return
        path = Path(selected)
        if path.suffix.casefold() != ".xlsx":
            path = path.with_suffix(".xlsx")
        try:
            with session_scope() as session:
                client_types = [category.name for category in categories.list_client_types(session)]
            _write_template(path, client_types)
        except Exception:
            self._show_error(ar.IMP_TEMPLATE_ERROR)
            return
        QMessageBox.information(self, ar.IMP_TEMPLATE_TITLE, ar.IMP_TEMPLATE_SAVED)

    def _sale_type_label(self, sale_type: str) -> str:
        """Translate internal sale type codes for the preview."""
        return _sale_type_label(sale_type)

    def _set_cell(self, row: int, column: int, text: str) -> QTableWidgetItem:
        """Place centered preview text and return the table item."""
        item = QTableWidgetItem(text)
        item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
        self.preview_table.setItem(row, column, item)
        return item

    def _show_error(self, message: str) -> None:
        """Show an import-page message in Arabic."""
        QMessageBox.warning(self, ar.IMP_PAGE_TITLE, message)


def _sheet_names(path: Path) -> list[str]:
    """List available worksheets for supported workbook formats."""
    if path.suffix.casefold() == ".xls":
        import pandas as pd

        book = pd.ExcelFile(path, engine="xlrd")
        try:
            return list(book.sheet_names)
        finally:
            book.close()
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _has_discrepancy(row: importer.ImportRow) -> bool:
    """Return whether any comparable spreadsheet value differs from calculation."""
    calculation = row.calculation
    if calculation is None:
        return False
    expected_monthly = calculation.monthly_list[0] if calculation.monthly_list else None
    compared_values = [
        sheet_value is not None and sheet_value != calculated_value
        for sheet_value, calculated_value in (
            (row.sheet_financed, calculation.financed),
            (row.sheet_profit, calculation.profit),
        )
    ]
    if row.sale_type == "installment":
        compared_values.append(
            row.sheet_monthly is not None and row.sheet_monthly != expected_monthly
        )
    return any(compared_values)


def _row_diagnostics(row: importer.ImportRow) -> tuple[str, str]:
    """Build the status and readable diagnostic text for a preview row."""
    diagnostics = tuple(row.errors) + tuple(row.warnings)
    if row.errors:
        return ar.IMP_STATUS_ERROR, "\n".join(diagnostics)
    if row.warnings:
        return ar.IMP_STATUS_WARNING, "\n".join(diagnostics)
    return ar.IMP_STATUS_OK, ar.IMP_STATUS_OK


def _format_integer(value: int | None) -> str:
    """Format an optional integer using western-digit thousands separators."""
    return f"{value:,}" if value is not None else "—"


def _format_date(value: date | None) -> str:
    """Format optional dates in the UI's documented day/month/year order."""
    return value.strftime("%d/%m/%Y") if value is not None else "—"


def _sale_type_label(sale_type: str) -> str:
    """Translate stored sale-type identifiers for page and report output."""
    return {
        "cash": ar.IMP_TYPE_CASH,
        "installment": ar.IMP_TYPE_INSTALLMENT,
        "credit": ar.IMP_TYPE_CREDIT,
    }.get(sale_type, sale_type)


def _write_template(path: Path, client_types: list[str] | None = None) -> None:
    """Create the sales sheet template (cash, installment and credit) with dropdowns.

    The required columns are blue; the optional ones (phone, client type,
    payment interval and the phone's colour, battery, IMEI and REF) are grey.
    """
    from app.i18n import is_rtl

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = importer.SHEET_NAME[:31]
    worksheet.sheet_view.rightToLeft = is_rtl()
    fields = tuple(importer.HEADERS)
    headers = [importer.HEADERS[field] for field in fields]
    worksheet.append(headers)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}500"
    for column, (field, header) in enumerate(zip(fields, headers), start=1):
        cell = worksheet.cell(row=1, column=column)
        cell.font = Font(bold=True, color="FFFFFF")
        optional = field in importer.OPTIONAL_HEADERS
        cell.fill = PatternFill(fill_type="solid", fgColor="5B6577" if optional else "0758CD")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if optional:
            cell.comment = Comment(ar.IMP_TEMPLATE_OPTIONAL_NOTE, "AkremMobile")
        worksheet.column_dimensions[get_column_letter(column)].width = max(
            17, min(34, len(header) + 4)
        )
    imei_letter = get_column_letter(fields.index("imei") + 1)
    worksheet.column_dimensions[imei_letter].number_format = "@"

    list_rules = [
        ("sale_type", ar.IMP_TEMPLATE_TYPES),
        ("rate", "0,5,10,15,20,25,30,35,40,45,50"),
    ]
    if client_types:
        choices = ",".join(name.replace(",", " ") for name in client_types)
        if len(choices) < 250:  # Excel's limit for an inline list
            list_rules.append(("client_type", choices))
    for field, choices in list_rules:
        letter = get_column_letter(fields.index(field) + 1)
        validation = DataValidation(type="list", formula1=f'"{choices}"', allow_blank=True)
        worksheet.add_data_validation(validation)
        validation.add(f"{letter}2:{letter}500")
    for field in ("months", "payment_interval"):
        letter = get_column_letter(fields.index(field) + 1)
        validation = DataValidation(type="whole", operator="between", formula1="1", formula2="60", allow_blank=True)
        worksheet.add_data_validation(validation)
        validation.add(f"{letter}2:{letter}500")
    workbook.save(path)


def _save_import_report(preview: importer.ImportPreview, result: importer.ImportResult) -> Path:
    """Write a summary and row diagnostics workbook under the exports directory."""
    config.ensure_data_dirs()
    export_dir = config.data_dir() / "exports"
    export_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    output = export_dir / f"import_report_{timestamp}.xlsx"

    workbook = Workbook()
    summary = workbook.active
    summary.title = ar.IMP_REPORT_SUMMARY[:31]
    summary.append((ar.IMP_REPORT_IMPORTED, result.imported))
    summary.append((ar.IMP_REPORT_SKIPPED, result.skipped))
    summary.append((ar.IMP_REPORT_WARNINGS, result.warnings))
    duplicate_rows = ", ".join(str(row) for row in result.duplicate_rows) or ar.IMP_NONE
    summary.append((ar.IMP_REPORT_DUPLICATE_ROWS, duplicate_rows))

    rows = workbook.create_sheet(ar.IMP_REPORT_ROWS[:31])
    rows.append(
        (
            ar.IMP_COLUMN_ROW,
            ar.IMP_COLUMN_CUSTOMER,
            ar.IMP_COLUMN_PRODUCT,
            ar.IMP_COLUMN_TYPE,
            ar.IMP_COLUMN_STATUS,
            ar.IMP_COLUMN_SHEET_FINANCED,
            ar.IMP_COLUMN_CALC_FINANCED,
            ar.IMP_COLUMN_SHEET_MONTHLY,
            ar.IMP_COLUMN_CALC_MONTHLY,
            ar.IMP_COLUMN_SHEET_PROFIT,
            ar.IMP_COLUMN_CALC_PROFIT,
        )
    )
    for row in preview.rows:
        calc = row.calculation
        sheet_monthly = row.sheet_monthly
        calc_monthly = calc.monthly_list[0] if calc and calc.monthly_list else None
        status, diagnostics = _row_diagnostics(row)
        rows.append(
            (
                row.source_row,
                row.full_name,
                row.product,
                _sale_type_label(row.sale_type) if row.sale_type else ar.IMP_CHOOSE_TYPE,
                f"{status}: {diagnostics}",
                row.sheet_financed,
                calc.financed if calc else None,
                sheet_monthly,
                calc_monthly,
                row.sheet_profit,
                calc.profit if calc else None,
            )
        )
    for worksheet in (summary, rows):
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions if worksheet is rows else "A1:B3"
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="0758CD")
        for column_cells in worksheet.columns:
            letter = column_cells[0].column_letter
            width = max(len(str(cell.value or "")) for cell in column_cells)
            worksheet.column_dimensions[letter].width = min(44, max(14, width + 2))
    workbook.save(output)
    return output
