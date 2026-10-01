"""Owner-managed product price book for sales entry.

A product needs only a name and a selling (cash) price; the wholesale price
is optional and stored as 0 when unknown. Each product also carries a default
installment plan (duration, months between payments, optional rate) that the
sale forms start from; every sale can still use its own plan. The catalog can also be filled
from an Excel sheet: :func:`write_product_template` creates the sheet,
:func:`preview_product_workbook` checks every row without touching the
database, and :func:`import_products` saves the valid rows.
"""

from __future__ import annotations

from dataclasses import dataclass, replace
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import importlib
from pathlib import Path
import re
from typing import Any, Iterable, NamedTuple

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.db.models import Product, Sale
from app.services import auth, calc
from app.services.spreadsheet import is_blank as _is_blank
from app.services.spreadsheet import normalize_header as _normalize_header
from app.services.spreadsheet import western_digits as _western_digits


class ProductNameTakenError(ValueError):
    """Another catalog product already has this name."""


DEFAULT_PLAN_MONTHS = 6


@dataclass(frozen=True)
class ProductPlan:
    """A product's default installment plan.

    ``interval`` is the number of months between payments (1 = monthly).
    ``rate`` None means "use the shop's rate preset for this duration".
    """

    months: int = DEFAULT_PLAN_MONTHS
    interval: int = 1
    rate: int | None = None

    def resolved_rate(self, presets: dict[int, int]) -> int | None:
        """The product's own rate, else the preset for its duration, else None."""
        return self.rate if self.rate is not None else presets.get(self.months)

    def payment_count(self) -> int:
        """Number of payments in this plan."""
        return len(calc.payment_offsets(self.months, self.interval))


class CatalogEntry(NamedTuple):
    """What a sale form needs to offer one product with its default plan."""

    name: str
    cash_price: int
    wholesale_price: int
    months: int = DEFAULT_PLAN_MONTHS
    interval: int = 1
    rate: int | None = None


def plan_of(product: Product) -> ProductPlan:
    """The default plan stored on a product."""
    return ProductPlan(
        months=product.default_months or DEFAULT_PLAN_MONTHS,
        interval=product.payment_interval or 1,
        rate=product.default_rate,
    )


def catalog_entries(session: Session, presets: dict[int, int]) -> list[CatalogEntry]:
    """Active products with their default plan and the rate that plan resolves to."""
    entries = []
    for product in list_products(session):
        plan = plan_of(product)
        entries.append(CatalogEntry(
            product.name, product.cash_price, product.wholesale_price,
            plan.months, plan.interval, plan.resolved_rate(presets),
        ))
    return entries


def _validate_plan(plan: ProductPlan) -> ProductPlan:
    """Check duration, interval and optional rate of a default plan."""
    calc.validate_plan(plan.months, plan.interval)
    if plan.rate is not None and (
        isinstance(plan.rate, bool) or not isinstance(plan.rate, int) or not 0 <= plan.rate <= 100
    ):
        raise ValueError("Plan rate must be a whole number between 0 and 100")
    return plan


def list_products(session: Session, *, include_inactive: bool = False) -> list[Product]:
    """Return catalog items alphabetically, hiding archived entries by default."""
    query = select(Product)
    if not include_inactive:
        query = query.where(Product.active.is_(True))
    return list(session.scalars(query.order_by(Product.name, Product.id)))


def create_product(
    session: Session,
    owner_user_id: int,
    *,
    name: str,
    cash_price: int,
    wholesale_price: int = 0,
    plan: ProductPlan | None = None,
) -> Product:
    """Create a catalog product after verifying the active owner account.

    Name and selling price are required; the wholesale price defaults to 0
    and the plan to 6 months paid monthly at the preset rate.
    """
    auth.require_owner(session, owner_user_id)
    clean_name = _validate_name(name)
    _validate_amount("wholesale price", wholesale_price)
    _validate_price(cash_price)
    clean_plan = _validate_plan(plan or ProductPlan())
    product = Product(
        name=clean_name,
        wholesale_price=wholesale_price,
        cash_price=cash_price,
        default_months=clean_plan.months,
        payment_interval=clean_plan.interval,
        default_rate=clean_plan.rate,
        active=True,
    )
    session.add(product)
    try:
        session.flush()
    except IntegrityError as error:
        raise ProductNameTakenError("A catalog product already uses this name") from error
    return product


def update_product(
    session: Session,
    owner_user_id: int,
    product_id: int,
    *,
    name: str,
    cash_price: int,
    wholesale_price: int | None = None,
    plan: ProductPlan | None = None,
) -> Product:
    """Update the active catalog price while old sales keep their stored snapshot.

    ``wholesale_price=None`` and ``plan=None`` keep the stored values.
    """
    auth.require_owner(session, owner_user_id)
    product = session.get(Product, product_id)
    if product is None:
        raise ValueError("Catalog product not found")
    product.name = _validate_name(name)
    if wholesale_price is not None:
        product.wholesale_price = _validate_amount("wholesale price", wholesale_price)
    product.cash_price = _validate_price(cash_price)
    if plan is not None:
        _set_plan(product, _validate_plan(plan))
    try:
        session.flush()
    except IntegrityError as error:
        raise ProductNameTakenError("A catalog product already uses this name") from error
    return product


def _set_plan(product: Product, plan: ProductPlan) -> None:
    product.default_months = plan.months
    product.payment_interval = plan.interval
    product.default_rate = plan.rate


def set_product_active(
    session: Session,
    owner_user_id: int,
    product_id: int,
    active: bool,
) -> Product:
    """Activate or archive a catalog item without changing historical sales."""
    auth.require_owner(session, owner_user_id)
    if not isinstance(active, bool):
        raise ValueError("Product active state must be boolean")
    product = session.get(Product, product_id)
    if product is None:
        raise ValueError("Catalog product not found")
    product.active = active
    session.flush()
    return product


def find_active_product(session: Session, name: str) -> Product | None:
    """Find an active catalog entry by its trimmed exact product name."""
    clean_name = name.strip() if isinstance(name, str) else ""
    if not clean_name:
        return None
    return session.scalar(
        select(Product).where(Product.name == clean_name, Product.active.is_(True))
    )


def _validate_name(name: str) -> str:
    """Trim a non-empty product name to the database length limit."""
    if not isinstance(name, str):
        raise ValueError("Product name is required")
    clean_name = name.strip()
    if not clean_name:
        raise ValueError("Product name cannot be empty")
    if len(clean_name) > 180:
        raise ValueError("Product name cannot exceed 180 characters")
    return clean_name


def _validate_amount(label: str, value: int) -> int:
    """Validate whole-dinar prices and return the normalized number."""
    if isinstance(value, bool) or not isinstance(value, int) or value < 0:
        raise ValueError(f"{label.capitalize()} must be a non-negative whole number")
    return value


def _validate_price(value: int) -> int:
    """The selling price is required: a whole number of dinars above zero."""
    _validate_amount("cash price", value)
    if value == 0:
        raise ValueError("Cash price is required")
    return value


def sales_by_product(session: Session) -> dict[str, tuple[int, date | None]]:
    """Return ``{product name: (units sold, last sale date)}`` from recorded sales.

    Sales keep a copy of the product name, so renamed catalog entries only
    count sales made under their current name.
    """
    rows = session.execute(
        select(Sale.product, func.count(Sale.id), func.max(Sale.purchase_date)).group_by(Sale.product)
    )
    return {name: (int(count), last_date) for name, count, last_date in rows}


# --------------------------------------------------------------- Excel import
MAX_IMPORT_ROWS = 5000
_HEADER_SCAN_ROWS = 10
_TEMPLATE_ROWS = 1000

# Header labels recognised in an uploaded sheet, besides the template's own
# labels in every supported language (see ``_header_aliases``).
_EXTRA_ALIASES: dict[str, tuple[str, ...]] = {
    "name": (
        "اسم المنتج", "المنتج", "المنتوج", "اسم المنتوج", "الاسم", "السلعة",
        "name", "product", "product name", "item", "model",
        "nom", "produit", "nom du produit", "article", "désignation", "modèle",
    ),
    "price": (
        "السعر", "سعر البيع", "سعر الكاش", "سعر البيع نقدا", "سعر ديطاي", "سعر التجزئة", "سعر المنتج",
        "price", "cash price", "selling price", "sale price", "retail price",
        "prix", "prix de vente", "prix comptant", "prix cash",
    ),
    "wholesale": (
        "سعر الجملة", "سعر الشراء", "سعر الشراء بالجملة", "الجملة",
        "wholesale", "wholesale price", "cost", "cost price", "purchase price",
        "prix de gros", "prix d'achat", "gros",
    ),
    "months": (
        "مدة التقسيط", "عدد الأشهر", "عدد أشهر التقسيط", "الأشهر", "المدة",
        "months", "duration", "installment months", "mois", "durée", "nombre de mois",
    ),
    "interval": (
        "الدفع كل", "يدفع كل", "كل كم شهر", "الدفع كل أشهر",
        "every", "pay every", "payment every", "interval", "payment interval",
        "tous les", "paiement tous les", "périodicité",
    ),
    "rate": (
        "النسبة", "نسبة التقسيط", "الفائدة %", "rate", "installment rate", "taux", "taux de crédit",
    ),
}
_TEMPLATE_KEYS = {
    "name": ("PROD_TPL_NAME", "SET_PRODUCT_NAME"),
    "price": ("PROD_TPL_PRICE", "SET_PRODUCT_CASH"),
    "wholesale": ("PROD_TPL_WHOLESALE", "SET_PRODUCT_WHOLESALE"),
    "months": ("PROD_TPL_MONTHS",),
    "interval": ("PROD_TPL_INTERVAL",),
    "rate": ("PROD_TPL_RATE",),
}

# Row problems, as codes the UI translates.
ERROR_NAME_MISSING = "name_missing"
ERROR_NAME_TOO_LONG = "name_too_long"
ERROR_PRICE_MISSING = "price_missing"
ERROR_PRICE_INVALID = "price_invalid"
ERROR_WHOLESALE_INVALID = "wholesale_invalid"
ERROR_DUPLICATE = "duplicate_in_file"
ERROR_MONTHS_INVALID = "months_invalid"
ERROR_INTERVAL_INVALID = "interval_invalid"
ERROR_RATE_INVALID = "rate_invalid"
WARNING_NO_WHOLESALE = "no_wholesale"
WARNING_BELOW_WHOLESALE = "below_wholesale"

ACTION_CREATE = "create"
ACTION_UPDATE = "update"
ACTION_RESTORE = "restore"
ACTION_UNCHANGED = "unchanged"
ACTION_INVALID = "invalid"


class ProductWorkbookError(ValueError):
    """The uploaded file cannot be read as a product sheet.

    ``code`` is ``unreadable``, ``no_header`` (no name/price columns found),
    ``empty`` or ``too_many_rows``.
    """

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class ProductImportRow:
    """One checked spreadsheet row; optional values are None when the cell was blank."""

    row_number: int
    name: str
    cash_price: int | None
    wholesale_price: int | None
    errors: tuple[str, ...]
    warnings: tuple[str, ...]
    action: str
    existing_id: int | None = None
    months: int | None = None
    interval: int | None = None
    rate: int | None = None

    def plan_over(self, base: ProductPlan) -> ProductPlan:
        """The plan after applying this row's filled plan cells to ``base``."""
        return ProductPlan(
            months=self.months if self.months is not None else base.months,
            interval=self.interval if self.interval is not None else base.interval,
            rate=self.rate if self.rate is not None else base.rate,
        )

    @property
    def valid(self) -> bool:
        """True when the row can be saved."""
        return not self.errors


@dataclass(frozen=True)
class ProductImportPreview:
    """Every non-empty row of an uploaded sheet with what importing it would do."""

    path: Path
    sheet_name: str
    rows: tuple[ProductImportRow, ...]

    def count(self, action: str) -> int:
        """Number of rows with the given action."""
        return sum(1 for row in self.rows if row.action == action)

    @property
    def valid_rows(self) -> tuple[ProductImportRow, ...]:
        """Rows without errors."""
        return tuple(row for row in self.rows if row.valid)

    @property
    def invalid_rows(self) -> tuple[ProductImportRow, ...]:
        """Rows that will be skipped."""
        return tuple(row for row in self.rows if not row.valid)

    @property
    def has_changes(self) -> bool:
        """True when importing would create, update or restore a product."""
        return any(row.action in (ACTION_CREATE, ACTION_UPDATE, ACTION_RESTORE) for row in self.rows)


@dataclass(frozen=True)
class ProductImportResult:
    """What an import saved."""

    created: int
    updated: int
    restored: int
    unchanged: int
    skipped: int

    @property
    def saved(self) -> int:
        """Products created, updated or restored."""
        return self.created + self.updated + self.restored


def write_product_template(path: str | Path) -> Path:
    """Write an empty product sheet plus an instructions sheet, in the UI language.

    Only the name and price columns are required; the wholesale price and
    the default plan (months, payment every N months, rate) may stay empty.
    """
    from app.i18n import ar, is_rtl

    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _sheet_title(ar.PROD_TPL_SHEET)
    sheet.sheet_view.rightToLeft = is_rtl()
    headers = (
        (f"{ar.PROD_TPL_NAME} *", "0758CD", ar.PROD_TPL_NAME_NOTE, 34),
        (f"{ar.PROD_TPL_PRICE} *", "0758CD", ar.PROD_TPL_PRICE_NOTE, 22),
        (f"{ar.PROD_TPL_WHOLESALE} ({ar.PROD_TPL_OPTIONAL})", "5B6577", ar.PROD_TPL_WHOLESALE_NOTE, 28),
        (f"{ar.PROD_TPL_MONTHS} ({ar.PROD_TPL_OPTIONAL})", "5B6577", ar.PROD_TPL_MONTHS_NOTE, 26),
        (f"{ar.PROD_TPL_INTERVAL} ({ar.PROD_TPL_OPTIONAL})", "5B6577", ar.PROD_TPL_INTERVAL_NOTE, 26),
        (f"{ar.PROD_TPL_RATE} ({ar.PROD_TPL_OPTIONAL})", "5B6577", ar.PROD_TPL_RATE_NOTE, 22),
    )
    for column, (label, color, note, width) in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor=color)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.comment = Comment(note, "AkremMobile")
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.row_dimensions[1].height = 24
    sheet.freeze_panes = "A2"
    for letter in ("B", "C"):
        sheet.column_dimensions[letter].number_format = "#,##0"
    price_rule = DataValidation(type="whole", operator="greaterThan", formula1="0", allow_blank=True)
    price_rule.error = ar.PROD_TPL_PRICE_NOTE
    price_rule.showErrorMessage = True
    sheet.add_data_validation(price_rule)
    price_rule.add(f"B2:B{_TEMPLATE_ROWS + 1}")
    wholesale_rule = DataValidation(
        type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True
    )
    wholesale_rule.error = ar.PROD_TPL_WHOLESALE_NOTE
    wholesale_rule.showErrorMessage = True
    sheet.add_data_validation(wholesale_rule)
    wholesale_rule.add(f"C2:C{_TEMPLATE_ROWS + 1}")
    for letter, low, high, note in (
        ("D", 1, calc.MAX_PLAN_MONTHS, ar.PROD_TPL_MONTHS_NOTE),
        ("E", 1, calc.MAX_PLAN_MONTHS, ar.PROD_TPL_INTERVAL_NOTE),
        ("F", 0, 100, ar.PROD_TPL_RATE_NOTE),
    ):
        rule = DataValidation(type="whole", operator="between", formula1=str(low), formula2=str(high), allow_blank=True)
        rule.error = note
        rule.showErrorMessage = True
        sheet.add_data_validation(rule)
        rule.add(f"{letter}2:{letter}{_TEMPLATE_ROWS + 1}")

    help_sheet = workbook.create_sheet(_sheet_title(ar.PROD_TPL_HELP_SHEET))
    help_sheet.sheet_view.rightToLeft = is_rtl()
    help_sheet.column_dimensions["A"].width = 90
    lines = (
        ar.PROD_TPL_HELP_TITLE,
        ar.PROD_TPL_HELP_REQUIRED,
        ar.PROD_TPL_HELP_OPTIONAL,
        ar.PROD_TPL_HELP_PLAN,
        ar.PROD_TPL_HELP_NUMBERS,
        ar.PROD_TPL_HELP_EXISTING,
        "",
        ar.PROD_TPL_HELP_EXAMPLE,
        "iPhone 13 — 110000 — 90000 — 10 — 2",
        "Redmi Note 13 — 32000",
        "Galaxy A05 — 21000 — 18000 — 3 — 1 — 20",
    )
    for index, line in enumerate(lines, start=1):
        cell = help_sheet.cell(row=index, column=1, value=line)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        if index == 1:
            cell.font = Font(bold=True, size=13, color="0758CD")
    workbook.active = 0
    workbook.save(destination)
    return destination


def preview_product_workbook(session: Session, path: str | Path) -> ProductImportPreview:
    """Read an uploaded product sheet and check every row; nothing is saved.

    The first sheet with a name column and a price column is used; the header
    row may be anywhere in its first rows. Amounts may use Arabic digits,
    thousands separators and the ``دج``/``DA`` suffix.
    """
    source = Path(path)
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except Exception as error:  # openpyxl raises many types for bad files
        raise ProductWorkbookError("unreadable", f"Cannot read {source.name}") from error
    try:
        found = _find_product_sheet(workbook)
        if found is None:
            raise ProductWorkbookError("no_header", "No product name and price columns found")
        sheet_name, header_row, columns, data = found
        if len(data) > MAX_IMPORT_ROWS:
            raise ProductWorkbookError("too_many_rows", f"More than {MAX_IMPORT_ROWS} rows")
    finally:
        workbook.close()

    existing = {_name_key(product.name): product for product in session.scalars(select(Product))}
    seen: set[str] = set()
    rows: list[ProductImportRow] = []
    for offset, values in enumerate(data, start=1):
        cells = {field: _cell(values, index) for field, index in columns.items()}
        if all(_is_blank(value) for value in cells.values()):
            continue
        rows.append(_check_row(header_row + offset, cells, existing, seen))
    if not rows:
        raise ProductWorkbookError("empty", "The sheet has no product rows")
    return ProductImportPreview(path=source, sheet_name=sheet_name, rows=tuple(rows))


def import_products(
    session: Session,
    owner_user_id: int,
    rows: Iterable[ProductImportRow],
) -> ProductImportResult:
    """Save the valid rows of a preview; owner only.

    New names become active products. A name already in the catalog (matched
    without regard to case or extra spaces) gets the new price, the new
    wholesale price when one was given, and is restored if it was archived.
    """
    auth.require_owner(session, owner_user_id)
    existing = {_name_key(product.name): product for product in session.scalars(select(Product))}
    created = updated = restored = unchanged = skipped = 0
    for row in rows:
        if not row.valid or row.cash_price is None:
            skipped += 1
            continue
        product = existing.get(_name_key(row.name))
        if product is None:
            product = create_product(
                session, owner_user_id,
                name=row.name, cash_price=row.cash_price, wholesale_price=row.wholesale_price or 0,
                plan=row.plan_over(ProductPlan()),
            )
            existing[_name_key(product.name)] = product
            created += 1
            continue
        action = _action_for(product, row)
        if action == ACTION_UNCHANGED:
            unchanged += 1
            continue
        product.cash_price = _validate_price(row.cash_price)
        if row.wholesale_price is not None:
            product.wholesale_price = _validate_amount("wholesale price", row.wholesale_price)
        _set_plan(product, _validate_plan(row.plan_over(plan_of(product))))
        if not product.active:
            product.active = True
            restored += 1
        else:
            updated += 1
    try:
        session.flush()
    except IntegrityError as error:
        raise ProductNameTakenError("A catalog product already uses this name") from error
    return ProductImportResult(created, updated, restored, unchanged, skipped)


def _check_row(
    row_number: int,
    cells: dict[str, Any],
    existing: dict[str, Product],
    seen: set[str],
) -> ProductImportRow:
    """Validate one row and decide whether it creates or updates a product."""
    errors: list[str] = []
    warnings: list[str] = []
    name = " ".join(_text(cells.get("name")).split())
    if not name:
        errors.append(ERROR_NAME_MISSING)
    elif len(name) > 180:
        errors.append(ERROR_NAME_TOO_LONG)
    price_cell = cells.get("price")
    cash_price = None
    if _is_blank(price_cell):
        errors.append(ERROR_PRICE_MISSING)
    else:
        cash_price = _parse_amount(price_cell)
        if cash_price is None or cash_price <= 0:
            errors.append(ERROR_PRICE_INVALID)
            cash_price = None
    wholesale_cell = cells.get("wholesale")
    wholesale_price = None
    if not _is_blank(wholesale_cell):
        wholesale_price = _parse_amount(wholesale_cell)
        if wholesale_price is None:
            errors.append(ERROR_WHOLESALE_INVALID)
    months = _optional_whole(cells.get("months"), 1, calc.MAX_PLAN_MONTHS, ERROR_MONTHS_INVALID, errors)
    interval = _optional_whole(cells.get("interval"), 1, calc.MAX_PLAN_MONTHS, ERROR_INTERVAL_INVALID, errors)
    rate_cell = cells.get("rate")
    rate = _optional_whole(
        rate_cell.replace("%", "") if isinstance(rate_cell, str) else rate_cell, 0, 100, ERROR_RATE_INVALID, errors
    )
    key = _name_key(name)
    if name and key in seen:
        errors.append(ERROR_DUPLICATE)
    elif name:
        seen.add(key)
    product = existing.get(key) if name else None
    if not errors and (months is not None or interval is not None):
        base = plan_of(product) if product is not None else ProductPlan()
        plan_months = months if months is not None else base.months
        plan_interval = interval if interval is not None else base.interval
        if plan_interval > plan_months:
            errors.append(ERROR_INTERVAL_INVALID)
    row = ProductImportRow(
        row_number=row_number,
        name=name,
        cash_price=cash_price,
        wholesale_price=wholesale_price,
        errors=(),
        warnings=(),
        action=ACTION_INVALID,
        existing_id=product.id if product is not None else None,
        months=months,
        interval=interval,
        rate=rate,
    )
    if errors:
        action = ACTION_INVALID
    elif product is None:
        action = ACTION_CREATE
    else:
        action = _action_for(product, row)
    known_wholesale = wholesale_price if wholesale_price is not None else (product.wholesale_price if product else 0)
    if not errors and not known_wholesale:
        warnings.append(WARNING_NO_WHOLESALE)
    elif not errors and cash_price is not None and cash_price < known_wholesale:
        warnings.append(WARNING_BELOW_WHOLESALE)
    return replace(row, errors=tuple(errors), warnings=tuple(warnings), action=action)


def _action_for(product: Product, row: ProductImportRow) -> str:
    """What saving this row would do to an existing product."""
    if not product.active:
        return ACTION_RESTORE
    same_wholesale = row.wholesale_price is None or row.wholesale_price == product.wholesale_price
    current_plan = plan_of(product)
    same_plan = row.plan_over(current_plan) == current_plan
    unchanged = row.cash_price == product.cash_price and same_wholesale and same_plan
    return ACTION_UNCHANGED if unchanged else ACTION_UPDATE


def _optional_whole(value: Any, low: int, high: int, error: str, errors: list[str]) -> int | None:
    """A blank cell gives None; anything else must be a whole number in ``low..high``."""
    if _is_blank(value):
        return None
    number = _parse_amount(value)
    if number is None or not low <= number <= high:
        errors.append(error)
        return None
    return number


def _find_product_sheet(
    workbook: Any,
) -> tuple[str, int, dict[str, int], list[tuple[Any, ...]]] | None:
    """Return (sheet name, header row number, column indexes, data rows) of the first product sheet."""
    aliases = _header_aliases()
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        for header_row in range(1, _HEADER_SCAN_ROWS + 1):
            values = next(rows, None)
            if values is None:
                break
            columns: dict[str, int] = {}
            for index, value in enumerate(values):
                field = aliases.get(_clean_header(value))
                if field is not None and field not in columns:
                    columns[field] = index
            if "name" in columns and "price" in columns:
                data: list[tuple[Any, ...]] = []
                for row in rows:
                    data.append(row)
                    if len(data) > MAX_IMPORT_ROWS + 1:
                        break
                return sheet.title, header_row, columns, data
    return None


def _header_aliases() -> dict[str, str]:
    """Map every normalised header label to its field."""
    aliases: dict[str, str] = {}
    for field, labels in _EXTRA_ALIASES.items():
        for label in labels:
            aliases.setdefault(_clean_header(label), field)
    for code in ("ar", "en", "fr"):
        module = importlib.import_module(f"app.i18n.{code}")
        for field, keys in _TEMPLATE_KEYS.items():
            for key in keys:
                label = getattr(module, key, None)
                if isinstance(label, str):
                    aliases[_clean_header(label)] = field
    return aliases


def _clean_header(value: Any) -> str:
    """Normalise a header cell: drop ``*``, bracketed notes, case and Arabic letter variants."""
    if _is_blank(value):
        return ""
    text = re.sub(r"[(\[（].*?[)\]）]", " ", str(value)).replace("*", " ").replace(":", " ")
    return " ".join(_normalize_header(text).split())


def _name_key(name: str) -> str:
    """Catalog name identity: case- and spacing-insensitive."""
    return " ".join(name.split()).casefold()


def _text(value: Any) -> str:
    """Cell text; whole numbers read from Excel lose their ``.0``."""
    if _is_blank(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_amount(value: Any) -> int | None:
    """Whole dinars from a number or text like ``١١٠٬٠٠٠ دج``; None when not a non-negative amount."""
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value if value >= 0 else None
    text = _western_digits(str(value)).casefold()
    for token in ("دج", "د.ج", "da", "dzd", " ", " ", " ", "٬", ","):
        text = text.replace(token, "")
    text = text.replace("٫", ".")
    try:
        number = Decimal(text).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError):
        return None
    if not number.is_finite() or number < 0:
        return None
    return int(number)


def _cell(values: tuple[Any, ...], index: int) -> Any:
    return values[index] if index < len(values) else None


def _sheet_title(title: str) -> str:
    """Excel sheet names: at most 31 characters and none of []:*?/\\."""
    clean = "".join(" " if char in '[]:*?/\\' else char for char in title)
    return clean[:31] or "Sheet"
