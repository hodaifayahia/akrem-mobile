"""Products & stock Excel sheet: template, upload (preview + import) and export.

The layout follows the shop's "Gros et Détail" workbook:

``المنتج | سعر الجملة | سعر الكاش | البطارية | اللون | IMEI | REF | ملاحظات | الكمية``
plus optional plan columns (months, pays every N months, rate).

Only the product name and the selling (cash) price are required.

Two kinds of sheet are understood:

* **Stock sheet** (any of battery / colour / IMEI / REF / note / quantity
  columns present): every row is one phone (or ``الكمية`` phones). The
  product is found by name (case and spacing ignored) or created, and each
  phone becomes a stock unit. Uploading the same sheet twice does not
  duplicate stock: a row with an IMEI or REF updates that unit, and rows
  without them are matched against identical unsold units already in stock.
* **Price list** (name / prices / plan only): one row per product; existing
  products get the new prices and plan, archived ones are restored.

Repeated header rows inside the data (the shop's sheet has one per batch),
blank rows and the side totals are skipped.
"""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, replace
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import importlib
from pathlib import Path
import re
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.db.models import Product, StockItem
from app.services import auth, calc, stock
from app.services.products import (
    ProductNameTakenError,
    ProductPlan,
    _set_plan,
    _validate_amount,
    _validate_plan,
    _validate_price,
    create_product,
    name_key,
    plan_of,
)
from app.services.spreadsheet import is_blank as _is_blank
from app.services.spreadsheet import normalize_header as _normalize_header
from app.services.spreadsheet import western_digits as _western_digits

MAX_IMPORT_ROWS = 5000
_HEADER_SCAN_ROWS = 10
_TEMPLATE_ROWS = 1000

# Header labels recognised in an uploaded sheet, besides the template's own
# labels in every supported language (see ``_header_aliases``).
_EXTRA_ALIASES: dict[str, tuple[str, ...]] = {
    "name": (
        "اسم المنتج", "المنتج", "المنتوج", "اسم المنتوج", "الاسم", "السلعة",
        "name", "product", "product name", "item", "model",
        "nom", "produit", "nom du produit", "article", "désignation", "modèle",
    ),
    "price": (
        "السعر", "سعر البيع", "سعر الكاش", "سعر البيع نقدا", "سعر ديطاي", "سعر التجزئة", "سعر المنتج",
        "price", "cash price", "selling price", "sale price", "retail price",
        "prix", "prix de vente", "prix comptant", "prix cash", "détail", "prix détail",
    ),
    "wholesale": (
        "سعر الجملة", "سعر الشراء", "سعر الشراء بالجملة", "الجملة",
        "wholesale", "wholesale price", "cost", "cost price", "purchase price",
        "prix de gros", "prix d'achat", "gros",
    ),
    "months": (
        "مدة التقسيط", "عدد الأشهر", "عدد أشهر التقسيط", "الأشهر", "المدة",
        "months", "duration", "installment months", "mois", "durée", "nombre de mois",
    ),
    "interval": (
        "الدفع كل", "يدفع كل", "كل كم شهر", "الدفع كل أشهر",
        "every", "pay every", "payment every", "interval", "payment interval",
        "tous les", "paiement tous les", "périodicité",
    ),
    "rate": (
        "النسبة", "نسبة التقسيط", "الفائدة %", "rate", "installment rate", "taux", "taux de crédit",
    ),
    "battery": ("البطارية", "صحة البطارية", "battery", "battery health", "batterie", "état batterie"),
    "color": ("اللون", "color", "colour", "couleur"),
    "imei": ("imei", "imei 1", "imei1", "satuts", "status", "statut", "serial", "رقم imei", "الرقم التسلسلي"),
    "reference": ("ref", "réf", "reference", "référence", "المرجع", "الرقم المرجعي", "مرجع"),
    "note": ("ملاحظات", "ملاحظة", "note", "notes", "remarque", "remarques", "comment", "comments"),
    "quantity": ("الكمية", "العدد", "quantity", "qty", "quantité", "qté"),
}
_TEMPLATE_KEYS = {
    "name": ("PROD_TPL_NAME", "SET_PRODUCT_NAME"),
    "price": ("PROD_TPL_PRICE", "SET_PRODUCT_CASH"),
    "wholesale": ("PROD_TPL_WHOLESALE", "SET_PRODUCT_WHOLESALE"),
    "months": ("PROD_TPL_MONTHS",),
    "interval": ("PROD_TPL_INTERVAL",),
    "rate": ("PROD_TPL_RATE",),
    "battery": ("PROD_TPL_BATTERY",),
    "color": ("PROD_TPL_COLOR",),
    "imei": ("PROD_TPL_IMEI",),
    "reference": ("PROD_TPL_REF",),
    "note": ("PROD_TPL_NOTE",),
    "quantity": ("PROD_TPL_QUANTITY",),
}
STOCK_FIELDS = frozenset({"battery", "color", "imei", "reference", "note", "quantity"})

# Row problems, as codes the UI translates.
ERROR_NAME_MISSING = "name_missing"
ERROR_NAME_TOO_LONG = "name_too_long"
ERROR_NAME_INVALID = "name_invalid"
ERROR_PRICE_MISSING = "price_missing"
ERROR_PRICE_INVALID = "price_invalid"
ERROR_WHOLESALE_INVALID = "wholesale_invalid"
ERROR_DUPLICATE = "duplicate_in_file"
ERROR_MONTHS_INVALID = "months_invalid"
ERROR_INTERVAL_INVALID = "interval_invalid"
ERROR_RATE_INVALID = "rate_invalid"
ERROR_QUANTITY_INVALID = "quantity_invalid"
ERROR_IMEI_DUPLICATE = "imei_duplicate"
WARNING_NO_WHOLESALE = "no_wholesale"
WARNING_BELOW_WHOLESALE = "below_wholesale"
WARNING_BATTERY_TEXT = "battery_text"
WARNING_STATUS_TEXT = "status_text"
WARNING_REF_DUPLICATE = "reference_duplicate"

# What a row does to the catalog.
ACTION_CREATE = "create"
ACTION_UPDATE = "update"
ACTION_RESTORE = "restore"
ACTION_UNCHANGED = "unchanged"
ACTION_EXISTING = "existing"  # stock row for a product already in the catalog
ACTION_INVALID = "invalid"

# What a stock row does to the stock.
UNIT_ADD = "unit_add"
UNIT_UPDATE = "unit_update"
UNIT_IN_STOCK = "unit_in_stock"
UNIT_SOLD = "unit_sold"


class ProductWorkbookError(ValueError):
    """The uploaded file cannot be read as a product sheet.

    ``code`` is ``unreadable``, ``no_header`` (no name/price columns found),
    ``empty`` or ``too_many_rows``.
    """

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class ProductImportRow:
    """One checked spreadsheet row; optional values are None when the cell was blank."""

    row_number: int
    name: str
    cash_price: int | None
    wholesale_price: int | None
    errors: tuple[str, ...]
    warnings: tuple[str, ...]
    action: str
    existing_id: int | None = None
    months: int | None = None
    interval: int | None = None
    rate: int | None = None
    # Stock sheet only.
    is_stock: bool = False
    quantity: int = 1
    color: str | None = None
    battery_health: int | None = None
    is_new: bool = False
    imei: str | None = None
    reference: str | None = None
    note: str | None = None
    unit_action: str | None = None
    units_to_add: int = 0
    existing_unit_id: int | None = None

    @property
    def valid(self) -> bool:
        """True when the row can be saved."""
        return not self.errors

    def plan_over(self, base: ProductPlan) -> ProductPlan:
        """The plan after applying this row's filled plan cells to ``base``."""
        return ProductPlan(
            months=self.months if self.months is not None else base.months,
            interval=self.interval if self.interval is not None else base.interval,
            rate=self.rate if self.rate is not None else base.rate,
        )

    def details(self) -> stock.UnitDetails:
        """The unit fields of a stock row."""
        return stock.UnitDetails(
            color=self.color, battery_health=self.battery_health, is_new=self.is_new,
            imei=self.imei, reference=self.reference, note=self.note,
        )


@dataclass(frozen=True)
class ProductImportPreview:
    """Every non-empty row of an uploaded sheet with what importing it would do."""

    path: Path
    sheet_name: str
    rows: tuple[ProductImportRow, ...]
    is_stock: bool = False

    def count(self, action: str) -> int:
        """Number of rows with the given catalog action."""
        return sum(1 for row in self.rows if row.action == action)

    @property
    def units_to_add(self) -> int:
        """Phones that would be added to stock."""
        return sum(row.units_to_add for row in self.rows if row.valid)

    def unit_count(self, unit_action: str) -> int:
        """Number of rows with the given stock action."""
        return sum(1 for row in self.rows if row.valid and row.unit_action == unit_action)

    @property
    def valid_rows(self) -> tuple[ProductImportRow, ...]:
        """Rows without errors."""
        return tuple(row for row in self.rows if row.valid)

    @property
    def invalid_rows(self) -> tuple[ProductImportRow, ...]:
        """Rows that will be skipped."""
        return tuple(row for row in self.rows if not row.valid)

    @property
    def has_changes(self) -> bool:
        """True when importing would change the catalog or the stock."""
        return any(
            row.valid and (
                row.action in (ACTION_CREATE, ACTION_UPDATE, ACTION_RESTORE)
                or row.unit_action in (UNIT_ADD, UNIT_UPDATE)
            )
            for row in self.rows
        )


@dataclass(frozen=True)
class ProductImportResult:
    """What an import saved."""

    created: int
    updated: int
    restored: int
    unchanged: int
    skipped: int
    units_added: int = 0
    units_updated: int = 0

    @property
    def saved(self) -> int:
        """Products created, updated or restored."""
        return self.created + self.updated + self.restored


# ----------------------------------------------------------------- template
def write_product_template(path: str | Path) -> Path:
    """Write an empty products & stock sheet plus an instructions sheet, in the UI language.

    Only the name and selling price columns are required (marked ``*``).
    """
    from app.i18n import ar, is_rtl

    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _sheet_title(ar.PROD_TPL_SHEET)
    sheet.sheet_view.rightToLeft = is_rtl()
    _write_header(sheet, template_columns())
    sheet.freeze_panes = "A2"
    for letter in ("B", "C"):
        sheet.column_dimensions[letter].number_format = "#,##0"
    _add_rule(sheet, "C", "greaterThan", 0, None, ar.PROD_TPL_PRICE_NOTE)
    _add_rule(sheet, "B", "greaterThanOrEqual", 0, None, ar.PROD_TPL_WHOLESALE_NOTE)
    _add_rule(sheet, "I", "between", 0, 500, ar.PROD_TPL_QUANTITY_NOTE)
    _add_rule(sheet, "J", "between", 1, calc.MAX_PLAN_MONTHS, ar.PROD_TPL_MONTHS_NOTE)
    _add_rule(sheet, "K", "between", 1, calc.MAX_PLAN_MONTHS, ar.PROD_TPL_INTERVAL_NOTE)
    _add_rule(sheet, "L", "between", 0, 100, ar.PROD_TPL_RATE_NOTE)
    sheet.column_dimensions["F"].number_format = "@"  # keep 15-digit IMEIs exact

    help_sheet = workbook.create_sheet(_sheet_title(ar.PROD_TPL_HELP_SHEET))
    help_sheet.sheet_view.rightToLeft = is_rtl()
    help_sheet.column_dimensions["A"].width = 100
    lines = (
        ar.PROD_TPL_HELP_TITLE,
        ar.PROD_TPL_HELP_REQUIRED,
        ar.PROD_TPL_HELP_STOCK,
        ar.PROD_TPL_HELP_BATTERY,
        ar.PROD_TPL_HELP_REF,
        ar.PROD_TPL_HELP_OPTIONAL,
        ar.PROD_TPL_HELP_PLAN,
        ar.PROD_TPL_HELP_NUMBERS,
        ar.PROD_TPL_HELP_EXISTING,
        "",
        ar.PROD_TPL_HELP_EXAMPLE,
        "IP 8 PLUS — 26000 — 30500 — 92% — Red — 356789012345678 — REF-00002",
        "SAM A07 4/64 — 22500 — 26000 — * — Black",
        "Redmi 15C 8/256 — 32000 — 35000",
    )
    for index, line in enumerate(lines, start=1):
        cell = help_sheet.cell(row=index, column=1, value=line)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        if index == 1:
            cell.font = Font(bold=True, size=13, color="0758CD")
    workbook.active = 0
    workbook.save(destination)
    return destination


def template_columns() -> list[tuple[str, bool, str, int]]:
    """``(header, required, note, width)`` for each template column, in sheet order."""
    from app.i18n import ar

    optional = f"({ar.PROD_TPL_OPTIONAL})"
    return [
        (f"{ar.PROD_TPL_NAME} *", True, ar.PROD_TPL_NAME_NOTE, 30),
        (f"{ar.PROD_TPL_WHOLESALE} {optional}", False, ar.PROD_TPL_WHOLESALE_NOTE, 22),
        (f"{ar.PROD_TPL_PRICE} *", True, ar.PROD_TPL_PRICE_NOTE, 18),
        (ar.PROD_TPL_BATTERY, False, ar.PROD_TPL_BATTERY_NOTE, 14),
        (ar.PROD_TPL_COLOR, False, ar.PROD_TPL_COLOR_NOTE, 14),
        (ar.PROD_TPL_IMEI, False, ar.PROD_TPL_IMEI_NOTE, 20),
        (ar.PROD_TPL_REF, False, ar.PROD_TPL_REF_NOTE, 14),
        (ar.PROD_TPL_NOTE, False, ar.PROD_TPL_NOTE_NOTE, 26),
        (ar.PROD_TPL_QUANTITY, False, ar.PROD_TPL_QUANTITY_NOTE, 10),
        (f"{ar.PROD_TPL_MONTHS} {optional}", False, ar.PROD_TPL_MONTHS_NOTE, 22),
        (f"{ar.PROD_TPL_INTERVAL} {optional}", False, ar.PROD_TPL_INTERVAL_NOTE, 22),
        (f"{ar.PROD_TPL_RATE} {optional}", False, ar.PROD_TPL_RATE_NOTE, 18),
    ]


def _write_header(sheet: Any, columns: list[tuple[str, bool, str, int]]) -> None:
    for column, (label, required, note, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=column, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="0758CD" if required else "5B6577")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if note:
            cell.comment = Comment(note, "AkremMobile")
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.row_dimensions[1].height = 24


def _add_rule(sheet: Any, letter: str, operator: str, low: int, high: int | None, note: str) -> None:
    rule = DataValidation(
        type="whole", operator=operator, formula1=str(low),
        formula2=None if high is None else str(high), allow_blank=True,
    )
    rule.error = note
    rule.showErrorMessage = True
    sheet.add_data_validation(rule)
    rule.add(f"{letter}2:{letter}{_TEMPLATE_ROWS + 1}")


# ------------------------------------------------------------------- export
def export_stock_workbook(session: Session, path: str | Path, units: Iterable[StockItem]) -> Path:
    """Write stock units in the template layout (re-importable) plus sale columns."""
    from app.i18n import ar, is_rtl

    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _sheet_title(ar.PROD_TPL_SHEET)
    sheet.sheet_view.rightToLeft = is_rtl()
    columns = template_columns()[:9] + [
        (ar.STOCK_COL_STATUS, False, "", 14),
        (ar.STOCK_COL_SOLD_TO, False, "", 24),
        (ar.STOCK_COL_SOLD_ON, False, "", 14),
    ]
    _write_header(sheet, columns)
    for unit in units:
        battery: object = "*" if unit.is_new else (unit.battery_health / 100 if unit.battery_health is not None else None)
        sale = unit.sale
        sheet.append([
            unit.product.name, unit.wholesale_price, unit.cash_price, battery, unit.color,
            unit.imei, unit.reference, unit.note, 1,
            ar.STOCK_STATUS_SOLD if unit.status == stock.STATUS_SOLD else ar.STOCK_STATUS_AVAILABLE,
            sale.customer.full_name if sale is not None else None,
            unit.sold_at,
        ])
    for row in sheet.iter_rows(min_row=2):
        row[1].number_format = row[2].number_format = "#,##0"
        if isinstance(row[3].value, float):
            row[3].number_format = "0%"
        row[5].number_format = "@"
        row[11].number_format = "dd/mm/yyyy"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(destination)
    return destination


# ------------------------------------------------------------------ preview
def preview_product_workbook(session: Session, path: str | Path) -> ProductImportPreview:
    """Read an uploaded products/stock sheet and check every row; nothing is saved.

    The first sheet with a name column and a price column is used; the header
    row may be anywhere in its first rows. Amounts may use Arabic digits,
    thousands separators and the ``دج``/``DA`` suffix.
    """
    source = Path(path)
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except Exception as error:  # openpyxl raises many types for bad files
        raise ProductWorkbookError("unreadable", f"Cannot read {source.name}") from error
    try:
        found = _find_product_sheet(workbook)
        if found is None:
            raise ProductWorkbookError("no_header", "No product name and price columns found")
        sheet_name, header_row, columns, data = found
        if len(data) > MAX_IMPORT_ROWS:
            raise ProductWorkbookError("too_many_rows", f"More than {MAX_IMPORT_ROWS} rows")
    finally:
        workbook.close()

    is_stock_sheet = bool(STOCK_FIELDS & columns.keys())
    aliases = _header_aliases()
    existing = {_name_key(product.name): product for product in session.scalars(select(Product))}
    seen_names: set[str] = set()
    seen_imeis: set[str] = set()
    seen_references: set[str] = set()
    rows: list[ProductImportRow] = []
    for offset, values in enumerate(data, start=1):
        cells = {field: _cell(values, index) for field, index in columns.items()}
        if all(_is_blank(value) for value in cells.values()) or _is_repeated_header(cells, aliases):
            continue
        rows.append(_check_row(
            header_row + offset, cells, existing, is_stock_sheet,
            seen_names, seen_imeis, seen_references,
        ))
    if not rows:
        raise ProductWorkbookError("empty", "The sheet has no product rows")
    if is_stock_sheet:
        rows = plan_stock(session, rows)
    return ProductImportPreview(path=source, sheet_name=sheet_name, rows=tuple(rows), is_stock=is_stock_sheet)


def _is_repeated_header(cells: dict[str, Any], aliases: dict[str, str]) -> bool:
    """The shop's sheet repeats its header row before each batch."""
    return aliases.get(_clean_header(cells.get("name"))) == "name"


def _check_row(
    row_number: int,
    cells: dict[str, Any],
    existing: dict[str, Product],
    is_stock_sheet: bool,
    seen_names: set[str],
    seen_imeis: set[str],
    seen_references: set[str],
) -> ProductImportRow:
    """Validate one row and decide what it does to the catalog."""
    errors: list[str] = []
    warnings: list[str] = []
    name = " ".join(_text(cells.get("name")).split())
    if not name:
        errors.append(ERROR_NAME_MISSING)
    elif name.startswith("="):
        errors.append(ERROR_NAME_INVALID)
    elif len(name) > 180:
        errors.append(ERROR_NAME_TOO_LONG)
    price_cell = cells.get("price")
    cash_price = None
    if _is_blank(price_cell):
        errors.append(ERROR_PRICE_MISSING)
    else:
        cash_price = _parse_amount(price_cell)
        if cash_price is None or cash_price <= 0:
            errors.append(ERROR_PRICE_INVALID)
            cash_price = None
    wholesale_cell = cells.get("wholesale")
    wholesale_price = None
    if not _is_blank(wholesale_cell):
        wholesale_price = _parse_amount(wholesale_cell)
        if wholesale_price is None:
            errors.append(ERROR_WHOLESALE_INVALID)
    months = _optional_whole(cells.get("months"), 1, calc.MAX_PLAN_MONTHS, ERROR_MONTHS_INVALID, errors)
    interval = _optional_whole(cells.get("interval"), 1, calc.MAX_PLAN_MONTHS, ERROR_INTERVAL_INVALID, errors)
    rate_cell = cells.get("rate")
    rate = _optional_whole(
        rate_cell.replace("%", "") if isinstance(rate_cell, str) else rate_cell, 0, 100, ERROR_RATE_INVALID, errors
    )
    key = _name_key(name)
    if name and not is_stock_sheet and key in seen_names:
        errors.append(ERROR_DUPLICATE)
    first_of_name = bool(name) and key not in seen_names
    if name:
        seen_names.add(key)
    product = existing.get(key) if name else None
    if not errors and (months is not None or interval is not None):
        base = plan_of(product) if product is not None else ProductPlan()
        plan_months = months if months is not None else base.months
        plan_interval = interval if interval is not None else base.interval
        if plan_interval > plan_months:
            errors.append(ERROR_INTERVAL_INVALID)

    unit: dict[str, Any] = {}
    if is_stock_sheet:
        unit = _check_unit_cells(cells, errors, warnings, seen_imeis, seen_references)

    row = ProductImportRow(
        row_number=row_number,
        name=name,
        cash_price=cash_price,
        wholesale_price=wholesale_price,
        errors=(),
        warnings=(),
        action=ACTION_INVALID,
        existing_id=product.id if product is not None else None,
        months=months,
        interval=interval,
        rate=rate,
        is_stock=is_stock_sheet,
        **unit,
    )
    if errors:
        action = ACTION_INVALID
    elif product is None:
        action = ACTION_CREATE if first_of_name else ACTION_EXISTING
    elif is_stock_sheet:
        action = _stock_catalog_action(product, row)
    else:
        action = _action_for(product, row)
    known_wholesale = wholesale_price if wholesale_price is not None else (product.wholesale_price if product else 0)
    if not errors and not known_wholesale:
        warnings.append(WARNING_NO_WHOLESALE)
    elif not errors and cash_price is not None and cash_price < known_wholesale:
        warnings.append(WARNING_BELOW_WHOLESALE)
    return replace(row, errors=tuple(errors), warnings=tuple(warnings), action=action)


def _check_unit_cells(
    cells: dict[str, Any],
    errors: list[str],
    warnings: list[str],
    seen_imeis: set[str],
    seen_references: set[str],
) -> dict[str, Any]:
    """Battery, colour, IMEI/status, REF, note and quantity of a stock row."""
    notes: list[str] = []
    is_new, health, battery_text = stock.parse_battery(cells.get("battery"))
    if battery_text:
        warnings.append(WARNING_BATTERY_TEXT)
        notes.append(battery_text)
    imei_cell = cells.get("imei")
    imei = stock.normalize_imei(imei_cell)
    if imei is None and not _is_blank(imei_cell):
        status_text = _text(imei_cell)
        if any(char.isalnum() for char in status_text):
            warnings.append(WARNING_STATUS_TEXT)
        notes.append(status_text)  # marks such as "+++++" are kept silently
    if imei is not None:
        if imei in seen_imeis:
            errors.append(ERROR_IMEI_DUPLICATE)
        seen_imeis.add(imei)
    reference = stock.normalize_reference(cells.get("reference"))
    if reference is not None:
        if reference in seen_references:
            warnings.append(WARNING_REF_DUPLICATE)
            reference = None
        else:
            seen_references.add(reference)
    note_cell = cells.get("note")
    if not _is_blank(note_cell):
        notes.insert(0, _text(note_cell))
    quantity = 1
    if not _is_blank(cells.get("quantity")):
        parsed = _parse_amount(cells.get("quantity"))
        if parsed is None or parsed > 500:
            errors.append(ERROR_QUANTITY_INVALID)
        else:
            quantity = parsed
    if quantity > 1 and (imei or reference):
        errors.append(ERROR_QUANTITY_INVALID)
    color = _text(cells.get("color")) or None
    return {
        "quantity": quantity,
        "color": " ".join(color.split())[:60] if color else None,
        "battery_health": health,
        "is_new": is_new,
        "imei": imei,
        "reference": reference,
        "note": " · ".join(notes)[:500] or None,
    }


def plan_stock(session: Session, rows: Iterable[ProductImportRow]) -> list[ProductImportRow]:
    """Decide what each valid stock row does to the stock, against the current database.

    A row with an IMEI or REF matching a unit updates that unit (or is skipped
    if it was sold). Other rows are matched against identical unsold units, so
    an already-uploaded sheet adds nothing the second time.
    """
    units = list(session.scalars(select(StockItem)))
    by_imei = {unit.imei: unit for unit in units if unit.imei}
    by_reference = {unit.reference: unit for unit in units if unit.reference}
    products = {product.id: product for product in session.scalars(select(Product))}
    product_by_key = {_name_key(product.name): product for product in products.values()}
    free: dict[tuple, int] = defaultdict(int)
    for unit in units:
        if unit.status == stock.STATUS_AVAILABLE and not unit.imei:
            product = products.get(unit.product_id)
            if product is not None:
                free[_unit_key(_name_key(product.name), unit.color, unit.battery_health, unit.is_new,
                               unit.cash_price, unit.wholesale_price)] += 1
    planned: list[ProductImportRow] = []
    for row in rows:
        if not row.valid or not row.is_stock:
            planned.append(row)
            continue
        match = by_imei.get(row.imei) if row.imei else None
        if match is None and row.reference:
            match = by_reference.get(row.reference)
        if match is not None:
            if match.status == stock.STATUS_SOLD:
                planned.append(replace(row, unit_action=UNIT_SOLD, units_to_add=0, existing_unit_id=match.id))
            else:
                changed = _unit_differs(match, row, product_by_key.get(_name_key(row.name)))
                planned.append(replace(
                    row, unit_action=UNIT_UPDATE if changed else UNIT_IN_STOCK,
                    units_to_add=0, existing_unit_id=match.id,
                ))
            continue
        if row.quantity == 0:
            planned.append(replace(row, unit_action=None, units_to_add=0))
            continue
        key = _unit_key(_name_key(row.name), row.color, row.battery_health, row.is_new,
                        row.cash_price, row.wholesale_price or 0)
        already = 0 if row.reference else min(free[key], row.quantity)
        free[key] -= already
        to_add = row.quantity - already
        planned.append(replace(row, unit_action=UNIT_ADD if to_add else UNIT_IN_STOCK, units_to_add=to_add))
    return planned


def _unit_key(name_key: str, color: str | None, health: int | None, is_new: bool, cash: int | None, wholesale: int) -> tuple:
    return (name_key, (color or "").casefold(), health, bool(is_new), cash, wholesale)


def _unit_differs(unit: StockItem, row: ProductImportRow, product: Product | None) -> bool:
    details = row.details().cleaned()
    return (
        (product is not None and unit.product_id != product.id)
        or unit.cash_price != row.cash_price
        or (row.wholesale_price is not None and unit.wholesale_price != row.wholesale_price)
        or (details.color or None) != (unit.color or None)
        or details.battery_health != unit.battery_health
        or details.is_new != unit.is_new
        or (details.imei is not None and details.imei != unit.imei)
        or (details.note is not None and details.note != unit.note)
    )


# ------------------------------------------------------------------- import
def import_products(
    session: Session,
    owner_user_id: int,
    rows: Iterable[ProductImportRow],
) -> ProductImportResult:
    """Save the valid rows of a preview; owner only, all or nothing (caller's transaction).

    Price list rows create or update products (matched by name, ignoring case
    and extra spaces) and restore archived ones. Stock rows create the product
    when it is new, then add, update or skip the unit as planned against the
    database at this moment.
    """
    auth.require_owner(session, owner_user_id)
    rows = list(rows)
    if any(row.is_stock for row in rows):
        rows = plan_stock(session, rows)
    existing = {_name_key(product.name): product for product in session.scalars(select(Product))}
    created = updated = restored = unchanged = skipped = units_added = units_updated = 0
    # Automatic references continue after the highest REF written in the sheet.
    references: set[str] = {row.reference for row in rows if row.valid and row.reference}
    for row in rows:
        if not row.valid or row.cash_price is None:
            skipped += 1
            continue
        key = _name_key(row.name)
        product = existing.get(key)
        if product is None:
            product = create_product(
                session, owner_user_id,
                name=row.name, cash_price=row.cash_price, wholesale_price=row.wholesale_price or 0,
                plan=row.plan_over(ProductPlan()),
            )
            existing[key] = product
            created += 1
        elif row.is_stock:
            action = _stock_catalog_action(product, row)
            if action == ACTION_UPDATE:
                _set_plan(product, _validate_plan(row.plan_over(plan_of(product))))
                updated += 1
            elif action == ACTION_RESTORE:
                product.active = True
                restored += 1
        else:
            action = _action_for(product, row)
            if action == ACTION_UNCHANGED:
                unchanged += 1
                continue
            product.cash_price = _validate_price(row.cash_price)
            if row.wholesale_price is not None:
                product.wholesale_price = _validate_amount("wholesale price", row.wholesale_price)
            _set_plan(product, _validate_plan(row.plan_over(plan_of(product))))
            if not product.active:
                product.active = True
                restored += 1
            else:
                updated += 1
        if not row.is_stock:
            continue
        if row.unit_action == UNIT_UPDATE and row.existing_unit_id is not None:
            unit = session.get(StockItem, row.existing_unit_id)
            details = row.details().cleaned()
            unit.product_id = product.id
            unit.cash_price = row.cash_price
            if row.wholesale_price is not None:
                unit.wholesale_price = row.wholesale_price
            unit.color = details.color
            unit.battery_health = details.battery_health
            unit.is_new = details.is_new
            unit.imei = details.imei or unit.imei
            unit.note = details.note if details.note is not None else unit.note
            units_updated += 1
        elif row.unit_action == UNIT_ADD:
            for _ in range(row.units_to_add):
                details = row.details().cleaned()
                stock._new_unit(  # noqa: SLF001 - same package, bulk path without per-unit owner checks
                    session, product, details, row.cash_price,
                    row.wholesale_price if row.wholesale_price is not None else product.wholesale_price,
                    taken=references,
                )
                units_added += 1
    try:
        session.flush()
    except IntegrityError as error:
        raise ProductNameTakenError("A catalog product or unit is duplicated") from error
    return ProductImportResult(created, updated, restored, unchanged, skipped, units_added, units_updated)


def _stock_catalog_action(product: Product, row: ProductImportRow) -> str:
    """Stock rows never change catalog prices; only a filled plan or an archived product."""
    if not product.active:
        return ACTION_RESTORE
    current = plan_of(product)
    return ACTION_UPDATE if row.plan_over(current) != current else ACTION_EXISTING


def _action_for(product: Product, row: ProductImportRow) -> str:
    """What saving this price-list row would do to an existing product."""
    if not product.active:
        return ACTION_RESTORE
    same_wholesale = row.wholesale_price is None or row.wholesale_price == product.wholesale_price
    current_plan = plan_of(product)
    same_plan = row.plan_over(current_plan) == current_plan
    unchanged = row.cash_price == product.cash_price and same_wholesale and same_plan
    return ACTION_UNCHANGED if unchanged else ACTION_UPDATE


# ------------------------------------------------------------------ helpers
def _optional_whole(value: Any, low: int, high: int, error: str, errors: list[str]) -> int | None:
    """A blank cell gives None; anything else must be a whole number in ``low..high``."""
    if _is_blank(value):
        return None
    number = _parse_amount(value)
    if number is None or not low <= number <= high:
        errors.append(error)
        return None
    return number


def _find_product_sheet(
    workbook: Any,
) -> tuple[str, int, dict[str, int], list[tuple[Any, ...]]] | None:
    """Return (sheet name, header row number, column indexes, data rows) of the first product sheet."""
    aliases = _header_aliases()
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        for header_row in range(1, _HEADER_SCAN_ROWS + 1):
            values = next(rows, None)
            if values is None:
                break
            columns: dict[str, int] = {}
            for index, value in enumerate(values):
                field = aliases.get(_clean_header(value))
                if field is not None and field not in columns:
                    columns[field] = index
            if "name" in columns and "price" in columns:
                data: list[tuple[Any, ...]] = []
                for row in rows:
                    data.append(row)
                    if len(data) > MAX_IMPORT_ROWS + 1:
                        break
                return sheet.title, header_row, columns, data
    return None


def _header_aliases() -> dict[str, str]:
    """Map every normalised header label to its field."""
    aliases: dict[str, str] = {}
    for field, labels in _EXTRA_ALIASES.items():
        for label in labels:
            aliases.setdefault(_clean_header(label), field)
    for code in ("ar", "en", "fr"):
        module = importlib.import_module(f"app.i18n.{code}")
        for field, keys in _TEMPLATE_KEYS.items():
            for key in keys:
                label = getattr(module, key, None)
                if isinstance(label, str):
                    aliases[_clean_header(label)] = field
    return aliases


def _clean_header(value: Any) -> str:
    """Normalise a header cell: drop ``*``, bracketed notes, case and Arabic letter variants."""
    if _is_blank(value) or not isinstance(value, str):
        return ""
    text = re.sub(r"[(\[（].*?[)\]）]", " ", value).replace("*", " ").replace(":", " ")
    return " ".join(_normalize_header(text).split())


def _name_key(name: str) -> str:
    """Catalog name identity: case- and spacing-insensitive."""
    return name_key(name)


def _text(value: Any) -> str:
    """Cell text; whole numbers read from Excel lose their ``.0``."""
    if _is_blank(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_amount(value: Any) -> int | None:
    """Whole dinars from a number or text like ``١١٠٬٠٠٠ دج``; None when not a non-negative amount."""
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value if value >= 0 else None
    text = _western_digits(str(value)).casefold()
    for token in ("دج", "د.ج", "da", "dzd", " ", " ", " ", "٬", ","):
        text = text.replace(token, "")
    text = text.replace("٫", ".")
    try:
        number = Decimal(text).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError):
        return None
    if not number.is_finite() or number < 0:
        return None
    return int(number)


def _cell(values: tuple[Any, ...], index: int) -> Any:
    return values[index] if index < len(values) else None


def _sheet_title(title: str) -> str:
    """Excel sheet names: at most 31 characters and none of []:*?/\\."""
    clean = "".join(" " if char in '[]:*?/\\' else char for char in title)
    return clean[:31] or "Sheet"
