"""Read-only monthly collection and owner profit reports."""

from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload, selectinload

from app.db.models import Installment, Sale, Setting
from app.services import auth
from app.services.status import Status, for_month


@dataclass(frozen=True, slots=True)
class CollectionRow:
    """One installment plus customer and sale details for a report."""

    customer: str
    phone: str
    product: str
    sale_type: str
    due_date: date
    amount_due: int
    amount_paid: int
    remaining: int
    status: Status


@dataclass(frozen=True, slots=True)
class OverdueRow:
    """One unpaid installment past its configured grace period."""

    customer: str
    phone: str
    product: str
    due_date: date
    amount_due: int
    amount_paid: int
    remaining: int
    days_late: int


@dataclass(frozen=True, slots=True)
class ProfitRow:
    """Sale totals grouped by the sale's purchase month."""

    year: int
    month: int
    sale_count: int
    wholesale: int
    total: int
    profit: int


def monthly_collection_rows(
    session: Session,
    *,
    year: int,
    month: int,
    today: date,
) -> list[CollectionRow]:
    """Return installment rows due in a selected month and their payment state."""
    if not 1 <= month <= 12:
        raise ValueError("month must be between 1 and 12")
    grace = _grace_days(session)
    installments = session.scalars(
        select(Installment)
        .options(
            selectinload(Installment.payments),
            joinedload(Installment.sale).joinedload(Sale.customer),
        )
        .where(
            Installment.due_date >= date(year, month, 1),
            Installment.due_date < _first_next_month(year, month),
        )
        .order_by(Installment.due_date, Installment.sale_id, Installment.installment_index)
    ).all()
    rows: list[CollectionRow] = []
    for installment in installments:
        sale = installment.sale
        customer = sale.customer
        paid = max(installment.amount_paid, sum(payment.amount for payment in installment.payments))
        rows.append(
            CollectionRow(
                customer.full_name,
                customer.phone or "",
                sale.product,
                sale.sale_type,
                installment.due_date,
                installment.amount_due,
                paid,
                max(0, installment.amount_due - paid),
                for_month([installment], year, month, today, grace),
            )
        )
    return rows


def overdue_rows(session: Session, *, today: date) -> list[OverdueRow]:
    """Return every installment that remains unpaid after due date and grace."""
    grace = _grace_days(session)
    installments = session.scalars(
        select(Installment)
        .options(
            selectinload(Installment.payments),
            joinedload(Installment.sale).joinedload(Sale.customer),
        )
        .where(Installment.due_date <= today - timedelta(days=grace + 1))
        .order_by(Installment.due_date, Installment.sale_id, Installment.installment_index)
    ).all()
    rows: list[OverdueRow] = []
    for installment in installments:
        paid = max(installment.amount_paid, sum(payment.amount for payment in installment.payments))
        remaining = max(0, installment.amount_due - paid)
        if remaining <= 0:
            continue
        sale = installment.sale
        customer = sale.customer
        rows.append(
            OverdueRow(
                customer.full_name,
                customer.phone or "",
                sale.product,
                installment.due_date,
                installment.amount_due,
                paid,
                remaining,
                (today - installment.due_date).days,
            )
        )
    return rows


def profit_by_month(
    session: Session,
    *,
    start_year: int,
    end_year: int,
    owner_user_id: int,
) -> list[ProfitRow]:
    """Aggregate sales and profit by purchase month for a verified owner."""
    auth.require_owner(session, owner_user_id)
    if start_year < 1 or end_year < start_year:
        raise ValueError("Invalid report year range")
    rows = session.scalars(
        select(Sale)
        .where(Sale.purchase_date >= date(start_year, 1, 1))
        .where(Sale.purchase_date < date(end_year + 1, 1, 1))
        .order_by(Sale.purchase_date, Sale.id)
    ).all()
    groups: dict[tuple[int, int], list[Sale]] = {}
    for sale in rows:
        groups.setdefault((sale.purchase_date.year, sale.purchase_date.month), []).append(sale)
    return [
        ProfitRow(
            year,
            month,
            len(sales),
            sum(sale.wholesale_price for sale in sales),
            sum(sale.total for sale in sales),
            sum(sale.profit for sale in sales),
        )
        for (year, month), sales in sorted(groups.items())
    ]


def export_collection_xlsx(rows: Iterable[CollectionRow], destination: str | Path) -> Path:
    """Write a monthly collection sheet with printable widths and Arabic headings."""
    from app.i18n import ar

    matrix = [[
        ar.REPORT_CUSTOMER, ar.REPORT_PHONE, ar.REPORT_PRODUCT, ar.REPORT_TYPE,
        ar.REPORT_DUE_DATE, ar.REPORT_DUE, ar.REPORT_PAID, ar.REPORT_REMAINING,
        ar.REPORT_STATUS,
    ]]
    labels = {
        Status.PAID: ar.CUST_STATUS_TIP_PAID,
        Status.FAILED: ar.CUST_STATUS_TIP_FAILED,
        Status.PENDING: ar.CUST_STATUS_TIP_PENDING,
        Status.NONE: ar.CUST_STATUS_TIP_NONE,
    }
    for row in rows:
        matrix.append([
            row.customer,
            row.phone,
            row.product,
            _sale_type_label(row.sale_type),
            row.due_date.strftime("%d/%m/%Y"),
            row.amount_due,
            row.amount_paid,
            row.remaining,
            labels[row.status],
        ])
    return _write_xlsx(matrix, destination, ar.REPORT_COLLECTION_TITLE)


def export_overdue_xlsx(rows: Iterable[OverdueRow], destination: str | Path) -> Path:
    """Write a list of outstanding past-due installments."""
    from app.i18n import ar

    matrix = [[
        ar.REPORT_CUSTOMER, ar.REPORT_PHONE, ar.REPORT_PRODUCT, ar.REPORT_DUE_DATE,
        ar.REPORT_DUE, ar.REPORT_PAID, ar.REPORT_REMAINING, ar.REPORT_DAYS_LATE,
    ]]
    matrix.extend([
        row.customer,
        row.phone,
        row.product,
        row.due_date.strftime("%d/%m/%Y"),
        row.amount_due,
        row.amount_paid,
        row.remaining,
        row.days_late,
    ] for row in rows)
    return _write_xlsx(matrix, destination, ar.REPORT_OVERDUE_TITLE)


def export_profit_xlsx(rows: Iterable[ProfitRow], destination: str | Path) -> Path:
    """Write an owner-only month-by-month sales/profit report."""
    from app.i18n import ar

    matrix = [[
        ar.REPORT_YEAR, ar.REPORT_MONTH, ar.REPORT_SALE_COUNT,
        ar.REPORT_WHOLESALE, ar.REPORT_TOTAL, ar.REPORT_PROFIT,
    ]]
    matrix.extend([
        row.year,
        f"{row.month:02d}",
        row.sale_count,
        row.wholesale,
        row.total,
        row.profit,
    ] for row in rows)
    return _write_xlsx(matrix, destination, ar.REPORT_PROFIT_TITLE)


def _write_xlsx(matrix: list[list[object]], destination: str | Path, title: str) -> Path:
    """Write a right-to-left workbook with a styled header and frozen first row."""
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = title[:31]
    worksheet.sheet_view.rightToLeft = True
    for row in matrix:
        worksheet.append(row)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="0758CD")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for column in worksheet.columns:
        values = [len(str(cell.value or "")) for cell in column]
        width = min(36, max(12, max(values, default=0) + 2))
        worksheet.column_dimensions[get_column_letter(column[0].column)].width = width
    workbook.save(path)
    return path


def _grace_days(session: Session) -> int:
    """Read the business grace period from settings with the five-day default."""
    setting = session.get(Setting, "grace_days")
    if setting is None:
        return 5
    try:
        value = json.loads(setting.value)
    except (TypeError, json.JSONDecodeError):
        return 5
    return value if isinstance(value, int) and not isinstance(value, bool) and value >= 0 else 5


def _first_next_month(year: int, month: int) -> date:
    """Return the first day after a selected month."""
    return date(year + 1, 1, 1) if month == 12 else date(year, month + 1, 1)


def _sale_type_label(value: str) -> str:
    """Translate a stored sale type for a report workbook."""
    from app.i18n import ar

    return {
        "cash": ar.SALE_CASH,
        "installment": ar.SALE_INSTALLMENT,
        "credit": ar.SALE_CREDIT,
    }.get(value, value)
